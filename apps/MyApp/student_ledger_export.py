"""
Executive Multi-Session Student Account Statement & Batchwise Financial Ledger Export Engine.
Generates publication-quality Portrait A4 PDF and multi-sheet OpenPyXL Excel (.xlsx) workbooks
matching official institutional branding and accounting standards of Sutawane Classes.
"""

import os
import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

from .export_theme import (
    COLOR_NAVY_PRIMARY,
    COLOR_NAVY_SECONDARY,
    COLOR_AMBER_ACCENT,
    COLOR_GREEN_PAID,
    COLOR_RED_PENDING,
    COLOR_BG_CARD,
    COLOR_BG_TOTALS,
    COLOR_BORDER,
    COLOR_SLATE_TEXT,
    COLOR_MUTED_TEXT,
    format_currency_inr,
    format_number_inr,
    FONT_FAMILY,
    FONT_FAMILY_BOLD,
)

# OpenPyXL Styles
THIN_BORDER = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1'),
)
TOTALS_BORDER = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='medium', color='1A2B4C'),
    bottom=Side(style='double', color='1A2B4C'),
)
HEADER_FILL = PatternFill(start_color="28375C", end_color="28375C", fill_type="solid")
SUBHEADER_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
TOTALS_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")


# =============================================================================
# Two-Pass Portrait A4 Running Canvas
# =============================================================================

class ExecutivePortraitCanvas(canvas.Canvas):
    """
    Two-pass canvas for ReportLab Portrait A4 PDFs.
    Draws professional running footer with 'Page X of Y' and institutional signature.
    """
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_footer(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_footer(self, page_count):
        self.saveState()
        self.setFont(FONT_FAMILY, 7.5)
        self.setFillColor(colors.HexColor(COLOR_MUTED_TEXT))
        page_w, page_h = A4

        # Left footer: Official statement signature
        self.drawString(20, 14, "Sutawane Classes Management Software • Confidential Official Account Statement")

        # Right footer: Page count
        self.drawRightString(page_w - 20, 14, f"Page {self._pageNumber} of {page_count}")

        # Thin divider line above footer
        self.setStrokeColor(colors.HexColor('#E2E8F0'))
        self.setLineWidth(0.5)
        self.line(20, 22, page_w - 20, 22)
        self.restoreState()


# =============================================================================
# Helper: Optimized Image Flowable
# =============================================================================

def get_optimized_image_flowable(img_path, max_w, max_h, fallback_text=None, title_style=None):
    """
    Creates a thumbnail in memory to ensure PDF stays lightweight and fast to render.
    """
    if img_path and os.path.exists(img_path):
        try:
            from PIL import Image as PILImage
            with PILImage.open(img_path) as im:
                im = im.convert('RGB')
                im.thumbnail((int(max_w * 3), int(max_h * 3)), PILImage.Resampling.LANCZOS)
                thumb_io = io.BytesIO()
                im.save(thumb_io, format='JPEG', quality=85, optimize=True)
                thumb_io.seek(0)
                return Image(thumb_io, width=max_w, height=max_h, kind='proportional')
        except Exception:
            try:
                return Image(img_path, width=max_w, height=max_h, kind='proportional')
            except Exception:
                pass
    if fallback_text and title_style:
        return Paragraph(fallback_text, title_style)
    return None


# =============================================================================
# 1. Executive PDF Export (Portrait A4)
# =============================================================================

def export_student_ledger_pdf(ledger_data):
    """
    Generates official publication-grade Portrait A4 PDF for student account statement:
    - Institutional branding header with master logo & contact info
    - Student photograph & bio identification strip
    - 4-Card lifetime financial summary ribbon
    - Session & batchwise breakdown with installment milestones & batch payments
    - Historical previous session dues table
    - All-time payment receipts audit trail
    - Official accountant stamp & parent signature blocks
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=20,
        rightMargin=20,
        topMargin=20,
        bottomMargin=30,
    )

    content_w = A4[0] - 40  # 555.27 pt

    styles = getSampleStyleSheet()

    # Typography styles
    title_style = ParagraphStyle(
        'LedgerTitle', parent=styles['Normal'],
        fontName=FONT_FAMILY_BOLD, fontSize=13, leading=16,
        textColor=colors.HexColor(COLOR_NAVY_PRIMARY)
    )
    subtitle_style = ParagraphStyle(
        'LedgerSubtitle', parent=styles['Normal'],
        fontName=FONT_FAMILY_BOLD, fontSize=9, leading=12,
        textColor=colors.HexColor(COLOR_NAVY_PRIMARY)
    )
    tagline_style = ParagraphStyle(
        'LedgerTagline', parent=styles['Normal'],
        fontName=FONT_FAMILY, fontSize=7, leading=9,
        textColor=colors.HexColor(COLOR_MUTED_TEXT)
    )
    meta_lbl_style = ParagraphStyle(
        'LedgerMetaLbl', parent=styles['Normal'],
        fontName=FONT_FAMILY_BOLD, fontSize=6.5, leading=8.5,
        textColor=colors.HexColor(COLOR_MUTED_TEXT)
    )
    meta_val_style = ParagraphStyle(
        'LedgerMetaVal', parent=styles['Normal'],
        fontName=FONT_FAMILY_BOLD, fontSize=7, leading=8.5,
        textColor=colors.HexColor(COLOR_SLATE_TEXT)
    )
    meta_val_red_style = ParagraphStyle(
        'LedgerMetaValRed', parent=styles['Normal'],
        fontName=FONT_FAMILY_BOLD, fontSize=7, leading=8.5,
        textColor=colors.HexColor(COLOR_RED_PENDING)
    )

    cell_bold_style = ParagraphStyle(
        'LedgerCellBold', parent=styles['Normal'],
        fontName=FONT_FAMILY_BOLD, fontSize=7.5, leading=9.5,
        textColor=colors.HexColor(COLOR_SLATE_TEXT)
    )
    cell_regular_style = ParagraphStyle(
        'LedgerCellReg', parent=styles['Normal'],
        fontName=FONT_FAMILY, fontSize=7.5, leading=9.5,
        textColor=colors.HexColor(COLOR_SLATE_TEXT)
    )
    cell_muted_style = ParagraphStyle(
        'LedgerCellMuted', parent=styles['Normal'],
        fontName=FONT_FAMILY, fontSize=6.5, leading=8.5,
        textColor=colors.HexColor(COLOR_MUTED_TEXT)
    )
    cell_green_bold = ParagraphStyle(
        'LedgerCellGreen', parent=styles['Normal'],
        fontName=FONT_FAMILY_BOLD, fontSize=7.5, leading=9.5,
        textColor=colors.HexColor(COLOR_GREEN_PAID), alignment=2
    )
    cell_red_bold = ParagraphStyle(
        'LedgerCellRed', parent=styles['Normal'],
        fontName=FONT_FAMILY_BOLD, fontSize=7.5, leading=9.5,
        textColor=colors.HexColor(COLOR_RED_PENDING), alignment=2
    )
    cell_right_bold = ParagraphStyle(
        'LedgerCellRightBold', parent=styles['Normal'],
        fontName=FONT_FAMILY_BOLD, fontSize=7.5, leading=9.5,
        textColor=colors.HexColor(COLOR_SLATE_TEXT), alignment=2
    )
    cell_right_reg = ParagraphStyle(
        'LedgerCellRightReg', parent=styles['Normal'],
        fontName=FONT_FAMILY, fontSize=7.5, leading=9.5,
        textColor=colors.HexColor(COLOR_SLATE_TEXT), alignment=2
    )
    th_style = ParagraphStyle(
        'LedgerTh', parent=styles['Normal'],
        fontName=FONT_FAMILY_BOLD, fontSize=7, leading=9,
        textColor=colors.white
    )
    th_right_style = ParagraphStyle(
        'LedgerThRight', parent=styles['Normal'],
        fontName=FONT_FAMILY_BOLD, fontSize=7, leading=9,
        textColor=colors.white, alignment=2
    )

    elements = []

    student = ledger_data['student']
    institute = ledger_data['institute']
    summary = ledger_data['lifetime_summary']
    sessions = ledger_data['sessions']
    previous_dues = ledger_data.get('previous_dues', [])
    payments_ledger = ledger_data.get('payments_ledger', [])

    # -------------------------------------------------------------
    # 1. Institutional Executive Header Block
    # -------------------------------------------------------------
    logo_path = institute.get('logo_path')
    logo_flowable = get_optimized_image_flowable(logo_path, max_w=44, max_h=44, fallback_text="<b>SC</b>", title_style=title_style) or Paragraph("<b>SC</b>", title_style)

    inst_name = (institute.get('name') or 'SUTAWANE CLASSES').upper()
    inst_tagline = institute.get('tagline') or 'Excellence in Coaching & Board Preparations'
    inst_contact = f"{institute.get('address', '')} • Phone: {institute.get('phone', '')}"

    left_header = [
        Paragraph(inst_name, title_style),
        Spacer(1, 1),
        Paragraph("OFFICIAL STUDENT ACCOUNT STATEMENT", subtitle_style),
        Spacer(1, 1),
        Paragraph(inst_tagline, tagline_style),
        Paragraph(inst_contact, cell_muted_style),
    ]

    badge_status = summary.get('status_label', 'VERIFIED')
    meta_rows = [
        [Paragraph("SCOPE:", meta_lbl_style), Paragraph("All Academic Sessions", meta_val_style)],
        [Paragraph("DATE:", meta_lbl_style), Paragraph(ledger_data.get('generated_at_str', ''), meta_val_style)],
        [Paragraph("STATUS:", meta_lbl_style), Paragraph(badge_status, meta_val_red_style if summary['pending_fees'] > 0 else meta_val_style)],
        [Paragraph("CLEARANCE:", meta_lbl_style), Paragraph(f"{summary['clearance_percentage']}%", meta_val_style)],
    ]
    meta_table = Table(meta_rows, colWidths=[55, 125])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor(COLOR_BG_CARD)),
        ('BOX', (0, 0), (-1, -1), 0.75, colors.HexColor(COLOR_BORDER)),
        ('TOPPADDING', (0, 0), (-1, -1), 2),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))

    hdr_table = Table([[logo_flowable, left_header, meta_table]], colWidths=[50, 325, 180])
    hdr_table.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 0),
        ('TOPPADDING', (0, 0), (-1, -1), 0),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
    ]))
    elements.append(hdr_table)
    elements.append(Spacer(1, 3))

    # Dual-tone accent line (100 pt Amber + remaining Navy)
    accent_bar = Table([['', '']], colWidths=[100, content_w - 100], rowHeights=[2.5])
    accent_bar.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, 0), colors.HexColor(COLOR_AMBER_ACCENT)),
        ('BACKGROUND', (1, 0), (1, 0), colors.HexColor(COLOR_NAVY_PRIMARY)),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 0),
        ('TOPPADDING', (0, 0), (-1, -1), 0),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
    ]))
    elements.append(accent_bar)
    elements.append(Spacer(1, 8))

    # -------------------------------------------------------------
    # 2. Student Bio & Photograph Identification Strip
    # -------------------------------------------------------------
    photo_disk_path = student.get('photo_disk_path')
    photo_fallback = f"<font size=14 color='#1A2B4C'><b>{student.get('initials', 'SC')}</b></font>"
    photo_flowable = get_optimized_image_flowable(photo_disk_path, max_w=54, max_h=54, fallback_text=photo_fallback, title_style=title_style) or Paragraph(photo_fallback, title_style)

    bio_col1 = [
        Paragraph(f"<b>{student.get('full_name', '')}</b>", ParagraphStyle('BName', parent=title_style, fontSize=11, leading=13)),
        Spacer(1, 1),
        Paragraph(f"Student Code: <b>{student.get('student_code', '')}</b>", cell_bold_style),
        Paragraph(f"Campus Branch: {student.get('primary_branch', 'Main Campus')}", cell_regular_style),
    ]

    bio_col2 = [
        Paragraph(f"Parent Mobile: <b>{student.get('parent_phone', '')}</b>", cell_regular_style),
        Paragraph(f"Student Mobile: {student.get('student_phone', '-')}", cell_regular_style),
        Paragraph(f"Address: {student.get('address', '-')}", cell_muted_style),
    ]

    bio_table = Table([[photo_flowable, bio_col1, bio_col2]], colWidths=[60, 265, 230])
    bio_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor(COLOR_BG_CARD)),
        ('BOX', (0, 0), (-1, -1), 0.75, colors.HexColor(COLOR_BORDER)),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(bio_table)
    elements.append(Spacer(1, 8))

    # -------------------------------------------------------------
    # 3. Lifetime Financial Ribbon (4 Metric KPI Cards)
    # -------------------------------------------------------------
    col_w = content_w / 4.0
    kpi_row = [
        [
            Paragraph("LIFETIME AGREED FEES", meta_lbl_style),
            Paragraph("TOTAL FEES PAID", meta_lbl_style),
            Paragraph("TOTAL PENDING BALANCE", meta_lbl_style),
            Paragraph("OVERALL CLEARANCE RATE", meta_lbl_style),
        ],
        [
            Paragraph(format_currency_inr(summary['agreed_fees']), cell_bold_style),
            Paragraph(format_currency_inr(summary['paid_fees']), cell_green_bold),
            Paragraph(format_currency_inr(summary['pending_fees']), cell_red_bold),
            Paragraph(f"{summary['clearance_percentage']}% ({summary['status_label']})", cell_bold_style),
        ]
    ]
    kpi_table = Table(kpi_row, colWidths=[col_w, col_w, col_w, col_w])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor(COLOR_BG_CARD)),
        ('BOX', (0, 0), (-1, -1), 0.75, colors.HexColor(COLOR_BORDER)),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor(COLOR_BORDER)),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    elements.append(kpi_table)
    elements.append(Spacer(1, 10))

    # -------------------------------------------------------------
    # 4. Academic Session & Batchwise Breakdown Sections
    # -------------------------------------------------------------
    sec_title = Paragraph("<b>SESSION & BATCHWISE FEE BREAKDOWN</b>", subtitle_style)
    elements.append(sec_title)
    elements.append(Spacer(1, 4))

    for session_grp in sessions:
        sess_name = session_grp.get('session_name', '')
        s_payable = format_currency_inr(session_grp.get('total_payable', 0))
        s_paid = format_currency_inr(session_grp.get('total_paid', 0))
        s_pend = format_currency_inr(session_grp.get('total_pending', 0))

        # Session Header Banner
        sess_hdr_text = f"<b>{sess_name.upper()}</b>  •  Agreed: {s_payable}  •  Collected: {s_paid}  •  Balance Due: {s_pend}"
        sess_hdr_para = Paragraph(sess_hdr_text, th_style)
        sess_bar = Table([[sess_hdr_para]], colWidths=[content_w], rowHeights=[18])
        sess_bar.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor(COLOR_NAVY_SECONDARY)),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 8),
            ('RIGHTPADDING', (0, 0), (-1, -1), 8),
            ('TOPPADDING', (0, 0), (-1, -1), 0),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 0),
        ]))

        session_flowables = [sess_bar, Spacer(1, 4)]

        for b_idx, batch in enumerate(session_grp.get('batches', []), start=1):
            # Batch subheader
            b_info = (
                f"<b>Batch {b_idx}: {batch['standard']}</b> (Roll: {batch['roll_number']}, Campus: {batch['branch_name']}) — "
                f"Net Fee: {format_currency_inr(batch['net_payable'])} | "
                f"Paid: {format_currency_inr(batch['paid_amount'])} | "
                f"Balance: {format_currency_inr(batch['pending_amount'])} [{batch['status']}]"
            )
            b_para = Paragraph(b_info, cell_bold_style)
            b_bar = Table([[b_para]], colWidths=[content_w])
            b_bar.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor(COLOR_BG_CARD)),
                ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor(COLOR_BORDER)),
                ('LEFTPADDING', (0, 0), (-1, -1), 6),
                ('TOPPADDING', (0, 0), (-1, -1), 3),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ]))
            session_flowables.append(b_bar)
            session_flowables.append(Spacer(1, 3))

            # Installments Table if available
            inst_list = batch.get('installments', [])
            if inst_list:
                inst_headers = [
                    Paragraph("#", th_style),
                    Paragraph("Milestone Label", th_style),
                    Paragraph("Due Date", th_style),
                    Paragraph("Scheduled", th_right_style),
                    Paragraph("Paid Amount", th_right_style),
                    Paragraph("Balance Due", th_right_style),
                    Paragraph("Status", th_style),
                ]
                inst_rows = [inst_headers]
                for inst in inst_list:
                    due_date_str = inst['due_date'].strftime('%d-%b-%Y') if inst.get('due_date') else '-'
                    inst_rows.append([
                        Paragraph(str(inst['installment_no']), cell_regular_style),
                        Paragraph(inst['label'], cell_regular_style),
                        Paragraph(due_date_str, cell_regular_style),
                        Paragraph(format_currency_inr(inst['amount_due']), cell_right_reg),
                        Paragraph(format_currency_inr(inst['paid_amount']), cell_green_bold),
                        Paragraph(format_currency_inr(inst['pending_amount']), cell_red_bold if inst['pending_amount'] > 0 else cell_right_reg),
                        Paragraph(inst['status'], cell_bold_style),
                    ])
                inst_table = Table(inst_rows, colWidths=[20, 150, 75, 80, 80, 80, 70])
                inst_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(COLOR_NAVY_PRIMARY)),
                    ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor(COLOR_BORDER)),
                    ('TOPPADDING', (0, 0), (-1, -1), 2.5),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
                    ('LEFTPADDING', (0, 0), (-1, -1), 4),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 4),
                    ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ]))
                session_flowables.append(inst_table)
                session_flowables.append(Spacer(1, 4))

        elements.extend(session_flowables)
        elements.append(Spacer(1, 6))

    # -------------------------------------------------------------
    # 5. Historical Previous Session Dues (if any)
    # -------------------------------------------------------------
    if previous_dues:
        elements.append(Paragraph("<b>HISTORICAL PREVIOUS SESSION DUES</b>", subtitle_style))
        elements.append(Spacer(1, 3))
        prev_headers = [
            Paragraph("Original Session", th_style),
            Paragraph("Carried To Session", th_style),
            Paragraph("Description", th_style),
            Paragraph("Outstanding", th_right_style),
            Paragraph("Collected", th_right_style),
            Paragraph("Balance", th_right_style),
            Paragraph("Status", th_style),
        ]
        prev_rows = [prev_headers]
        for due in previous_dues:
            prev_rows.append([
                Paragraph(due['original_session'], cell_regular_style),
                Paragraph(due['current_session'], cell_regular_style),
                Paragraph(due['description'], cell_muted_style),
                Paragraph(format_currency_inr(due['outstanding_amount']), cell_right_reg),
                Paragraph(format_currency_inr(due['collected_amount']), cell_green_bold),
                Paragraph(format_currency_inr(due['pending_amount']), cell_red_bold if due['pending_amount'] > 0 else cell_right_reg),
                Paragraph(due['status'], cell_bold_style),
            ])
        prev_table = Table(prev_rows, colWidths=[90, 90, 135, 60, 60, 60, 60])
        prev_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(COLOR_NAVY_PRIMARY)),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor(COLOR_BORDER)),
            ('TOPPADDING', (0, 0), (-1, -1), 2.5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        elements.append(prev_table)
        elements.append(Spacer(1, 8))

    # -------------------------------------------------------------
    # 6. All-Time Payment Receipts Ledger (Audit Trail)
    # -------------------------------------------------------------
    elements.append(Paragraph("<b>COMPLETE PAYMENT RECEIPTS HISTORY (ALL-TIME AUDIT TRAIL)</b>", subtitle_style))
    elements.append(Spacer(1, 4))

    ledger_headers = [
        Paragraph("#", th_style),
        Paragraph("Receipt #", th_style),
        Paragraph("Date", th_style),
        Paragraph("Session", th_style),
        Paragraph("Course / Standard", th_style),
        Paragraph("Mode", th_style),
        Paragraph("Received By", th_style),
        Paragraph("Amount", th_right_style),
        Paragraph("Running Total", th_right_style),
    ]
    ledger_rows = [ledger_headers]
    for p in payments_ledger:
        p_date_str = p['payment_date'].strftime('%d-%b-%Y') if p.get('payment_date') else '-'
        rcv_staff = p.get('received_by') or 'Staff'
        ledger_rows.append([
            Paragraph(str(p['index']), cell_regular_style),
            Paragraph(p['receipt_number'], cell_bold_style),
            Paragraph(p_date_str, cell_regular_style),
            Paragraph(p['session_name'], cell_regular_style),
            Paragraph(p['standard'], cell_regular_style),
            Paragraph(p['payment_mode'], cell_regular_style),
            Paragraph(rcv_staff, cell_regular_style),
            Paragraph(format_currency_inr(p['amount']), cell_green_bold),
            Paragraph(format_currency_inr(p['running_total']), cell_right_bold),
        ])

    # Grand Totals Row
    ledger_rows.append([
        Paragraph("<b>GRAND TOTAL VERIFIED COLLECTIONS:</b>", cell_bold_style),
        "", "", "", "", "", "",
        Paragraph(format_currency_inr(summary['paid_fees']), cell_green_bold),
        Paragraph(format_currency_inr(summary['paid_fees']), cell_right_bold),
    ])

    rec_table = Table(ledger_rows, colWidths=[20, 95, 65, 65, 85, 45, 60, 60, 60])
    rec_style = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(COLOR_NAVY_PRIMARY)),
        ('GRID', (0, 0), (-1, -2), 0.5, colors.HexColor(COLOR_BORDER)),
        ('TOPPADDING', (0, 0), (-1, -1), 2.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        # Totals row
        ('SPAN', (0, -1), (6, -1)),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor(COLOR_BG_TOTALS)),
        ('LINEABOVE', (0, -1), (-1, -1), 1.0, colors.HexColor(COLOR_NAVY_PRIMARY)),
        ('LINEBELOW', (0, -1), (-1, -1), 1.5, colors.HexColor(COLOR_NAVY_PRIMARY)),
    ]
    rec_table.setStyle(TableStyle(rec_style))
    elements.append(rec_table)
    elements.append(Spacer(1, 14))

    # -------------------------------------------------------------
    # 7. Official Verification & Signatures Block
    # -------------------------------------------------------------
    sig_note = Paragraph(
        f"<i>{institute.get('receipt_footer', '')} • This document is an official confidential account statement "
        f"generated by Sutawane Classes Management Software on {ledger_data.get('generated_at_str', '')}.</i>",
        cell_muted_style
    )
    sig_box_w = content_w / 3.0
    sig_rows = [
        [
            Paragraph("____________________________<br/><b>Student / Parent Signature</b><br/><font size=6 color='#64748B'>Acknowledged & Verified</font>", cell_regular_style),
            Paragraph("____________________________<br/><b>Branch Accountant Stamp</b><br/><font size=6 color='#64748B'>Account Verified</font>", cell_regular_style),
            Paragraph(f"____________________________<br/><b>Authorized Signatory</b><br/><font size=6 color='#64748B'>{institute.get('name', 'Sutawane Classes')}</font>", cell_regular_style),
        ]
    ]
    sig_table = Table(sig_rows, colWidths=[sig_box_w, sig_box_w, sig_box_w])
    sig_table.setStyle(TableStyle([
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))

    sig_container = KeepTogether([
        sig_note,
        Spacer(1, 10),
        sig_table,
    ])
    elements.append(sig_container)

    # Build PDF with two-pass portrait canvas
    doc.build(elements, canvasmaker=ExecutivePortraitCanvas)
    pdf_data = buffer.getvalue()
    buffer.close()
    return pdf_data


# =============================================================================
# 2. Executive Multi-Sheet Excel Export (.xlsx)
# =============================================================================

def export_student_ledger_excel(ledger_data):
    """
    Exports comprehensive student account statement to a 3-sheet Excel workbook:
    - Sheet 1: Account Statement (Bio, Lifetime KPIs, Session & Batchwise Summary)
    - Sheet 2: Installment Milestones (Batchwise schedules & statuses)
    - Sheet 3: Payment Receipts History (All-time chronological transactions with totals)
    """
    student = ledger_data['student']
    institute = ledger_data['institute']
    summary = ledger_data['lifetime_summary']
    sessions = ledger_data['sessions']
    previous_dues = ledger_data.get('previous_dues', [])
    payments_ledger = ledger_data.get('payments_ledger', [])
    gen_at = ledger_data.get('generated_at_str', '')

    wb = Workbook()

    bold_font = Font(name="Calibri", size=10, bold=True)
    normal_font = Font(name="Calibri", size=10)
    paid_font = Font(name="Calibri", size=10, bold=True, color="16A34A")
    due_font = Font(name="Calibri", size=10, bold=True, color="DC2626")
    totals_bold = Font(name="Calibri", size=11, bold=True, color="1A2B4C")
    totals_paid = Font(name="Calibri", size=11, bold=True, color="16A34A")
    totals_due = Font(name="Calibri", size=11, bold=True, color="DC2626")
    hdr_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")

    # -------------------------------------------------------------
    # Sheet 1: Account Statement
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Account Statement"
    ws1.views.sheetView[0].showGridLines = True

    # Title & Metadata
    ws1.merge_cells("A1:I1")
    ws1["A1"] = f"{institute.get('name', 'SUTAWANE CLASSES').upper()} — OFFICIAL STUDENT ACCOUNT STATEMENT"
    ws1["A1"].font = Font(name="Calibri", size=13, bold=True, color="1A2B4C")
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 26

    ws1.merge_cells("A2:I2")
    ws1["A2"] = f"Generated: {gen_at} • Scope: All Academic Sessions • Student: {student['full_name']} ({student['student_code']})"
    ws1["A2"].font = Font(name="Calibri", size=9.5, italic=True, color="64748B")
    ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 18

    # Student Bio Section
    ws1["A4"] = "Student Code:"
    ws1["B4"] = student.get('student_code', '')
    ws1["D4"] = "Parent Mobile:"
    ws1["E4"] = student.get('parent_phone', '')
    ws1["G4"] = "Primary Campus:"
    ws1["H4"] = student.get('primary_branch', '')

    ws1["A5"] = "Full Name:"
    ws1["B5"] = student.get('full_name', '')
    ws1["D5"] = "Student Mobile:"
    ws1["E5"] = student.get('student_phone', '-')
    ws1["G5"] = "Clearance Rate:"
    ws1["H5"] = f"{summary['clearance_percentage']}% ({summary['status_label']})"

    for r in [4, 5]:
        for c in [1, 4, 7]:
            ws1.cell(row=r, column=c).font = bold_font
            ws1.cell(row=r, column=c).fill = SUBHEADER_FILL
        for c in [2, 5, 8]:
            ws1.cell(row=r, column=c).font = normal_font
        ws1.row_dimensions[r].height = 20

    # Lifetime Financial Ribbon
    ws1["A7"] = "LIFETIME AGREED FEES"
    ws1["C7"] = "TOTAL FEES PAID"
    ws1["E7"] = "TOTAL PENDING BALANCE"
    ws1["G7"] = "OVERALL CLEARANCE"

    ws1["A8"] = float(summary['agreed_fees'])
    ws1["C8"] = float(summary['paid_fees'])
    ws1["E8"] = float(summary['pending_fees'])
    ws1["G8"] = f"{summary['clearance_percentage']}%"

    ws1["A8"].number_format = '₹#,##0.00'
    ws1["C8"].number_format = '₹#,##0.00'
    ws1["E8"].number_format = '₹#,##0.00'

    for col_idx in [1, 3, 5, 7]:
        ws1.cell(row=7, column=col_idx).font = Font(name="Calibri", size=9, bold=True, color="64748B")
        ws1.cell(row=7, column=col_idx).fill = SUBHEADER_FILL
        ws1.cell(row=7, column=col_idx).alignment = Alignment(horizontal="center")
        ws1.cell(row=8, column=col_idx).alignment = Alignment(horizontal="center")
    
    ws1["A8"].font = totals_bold
    ws1["C8"].font = totals_paid
    ws1["E8"].font = totals_due
    ws1["G8"].font = totals_bold

    # Table: Session & Batchwise Breakdown
    headers1 = [
        "#", "Academic Session", "Course / Batch", "Roll Number", "Campus Branch",
        "Net Agreed Fee (Rs.)", "Total Paid (Rs.)", "Pending Balance (Rs.)", "Clearance (%)", "Status"
    ]
    cur_row = 10
    ws1.row_dimensions[cur_row].height = 24
    for c_idx, h in enumerate(headers1, start=1):
        cell = ws1.cell(row=cur_row, column=c_idx, value=h)
        cell.font = hdr_font
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="right" if "Rs." in h or "(%)" in h else "left", vertical="center")
        cell.border = THIN_BORDER

    item_idx = 1
    start_data_row = cur_row + 1
    cur_row += 1

    for sess in sessions:
        for b in sess.get('batches', []):
            ws1.row_dimensions[cur_row].height = 20
            ws1.cell(row=cur_row, column=1, value=item_idx).alignment = Alignment(horizontal="center")
            ws1.cell(row=cur_row, column=2, value=sess['session_name'])
            ws1.cell(row=cur_row, column=3, value=b['standard'])
            ws1.cell(row=cur_row, column=4, value=b['roll_number'])
            ws1.cell(row=cur_row, column=5, value=b['branch_name'])

            c6 = ws1.cell(row=cur_row, column=6, value=float(b['net_payable']))
            c6.number_format = '#,##0.00'
            c6.alignment = Alignment(horizontal="right")

            c7 = ws1.cell(row=cur_row, column=7, value=float(b['paid_amount']))
            c7.number_format = '#,##0.00'
            c7.font = paid_font
            c7.alignment = Alignment(horizontal="right")

            c8 = ws1.cell(row=cur_row, column=8, value=float(b['pending_amount']))
            c8.number_format = '#,##0.00'
            c8.font = due_font if b['pending_amount'] > 0 else normal_font
            c8.alignment = Alignment(horizontal="right")

            ws1.cell(row=cur_row, column=9, value=f"{b['clearance_percentage']}%").alignment = Alignment(horizontal="right")
            ws1.cell(row=cur_row, column=10, value=b['status']).alignment = Alignment(horizontal="center")

            for col in range(1, 11):
                ws1.cell(row=cur_row, column=col).border = THIN_BORDER
            cur_row += 1
            item_idx += 1

    # Totals Row
    ws1.row_dimensions[cur_row].height = 24
    ws1.merge_cells(f"A{cur_row}:E{cur_row}")
    ws1[f"A{cur_row}"] = "TOTALS ACROSS ALL BATCHES:"
    ws1[f"A{cur_row}"].font = totals_bold
    ws1[f"A{cur_row}"].alignment = Alignment(horizontal="right", vertical="center")

    t6 = ws1.cell(row=cur_row, column=6, value=f"=SUM(F{start_data_row}:F{cur_row-1})")
    t6.number_format = '#,##0.00'
    t6.font = totals_bold
    t6.alignment = Alignment(horizontal="right", vertical="center")

    t7 = ws1.cell(row=cur_row, column=7, value=f"=SUM(G{start_data_row}:G{cur_row-1})")
    t7.number_format = '#,##0.00'
    t7.font = totals_paid
    t7.alignment = Alignment(horizontal="right", vertical="center")

    t8 = ws1.cell(row=cur_row, column=8, value=f"=SUM(H{start_data_row}:H{cur_row-1})")
    t8.number_format = '#,##0.00'
    t8.font = totals_due
    t8.alignment = Alignment(horizontal="right", vertical="center")

    ws1.cell(row=cur_row, column=9, value=f"{summary['clearance_percentage']}%").font = totals_bold
    ws1.cell(row=cur_row, column=9).alignment = Alignment(horizontal="right", vertical="center")

    ws1.cell(row=cur_row, column=10, value=summary['status_label']).font = totals_bold
    ws1.cell(row=cur_row, column=10).alignment = Alignment(horizontal="center", vertical="center")

    for col in range(1, 11):
        cell = ws1.cell(row=cur_row, column=col)
        cell.border = TOTALS_BORDER
        cell.fill = TOTALS_FILL

    # Auto-width
    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # -------------------------------------------------------------
    # Sheet 2: Installment Milestones
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Installment Milestones")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:J1")
    ws2["A1"] = f"{institute.get('name', 'SUTAWANE CLASSES').upper()} — BATCHWISE INSTALLMENT SCHEDULE & MILESTONES"
    ws2["A1"].font = Font(name="Calibri", size=13, bold=True, color="1A2B4C")
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 26

    ws2.merge_cells("A2:J2")
    ws2["A2"] = f"Student: {student['full_name']} ({student['student_code']}) • Scope: All Sessions"
    ws2["A2"].font = Font(name="Calibri", size=9.5, italic=True, color="64748B")
    ws2["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[2].height = 18

    headers2 = [
        "#", "Academic Session", "Course / Batch", "Roll No", "Milestone Label",
        "Due Date", "Amount Due (Rs.)", "Paid Amount (Rs.)", "Pending Balance (Rs.)", "Status"
    ]
    ws2.row_dimensions[4].height = 24
    for c_idx, h in enumerate(headers2, start=1):
        cell = ws2.cell(row=4, column=c_idx, value=h)
        cell.font = hdr_font
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="right" if "Rs." in h else "left", vertical="center")
        cell.border = THIN_BORDER

    cur_row2 = 5
    m_idx = 1
    start_row2 = 5

    for sess in sessions:
        for b in sess.get('batches', []):
            for inst in b.get('installments', []):
                ws2.row_dimensions[cur_row2].height = 20
                ws2.cell(row=cur_row2, column=1, value=m_idx).alignment = Alignment(horizontal="center")
                ws2.cell(row=cur_row2, column=2, value=sess['session_name'])
                ws2.cell(row=cur_row2, column=3, value=b['standard'])
                ws2.cell(row=cur_row2, column=4, value=b['roll_number'])
                ws2.cell(row=cur_row2, column=5, value=inst['label'])

                due_d_str = inst['due_date'].strftime('%d-%b-%Y') if inst.get('due_date') else '-'
                ws2.cell(row=cur_row2, column=6, value=due_d_str).alignment = Alignment(horizontal="center")

                c7 = ws2.cell(row=cur_row2, column=7, value=float(inst['amount_due']))
                c7.number_format = '#,##0.00'
                c7.alignment = Alignment(horizontal="right")

                c8 = ws2.cell(row=cur_row2, column=8, value=float(inst['paid_amount']))
                c8.number_format = '#,##0.00'
                c8.font = paid_font
                c8.alignment = Alignment(horizontal="right")

                c9 = ws2.cell(row=cur_row2, column=9, value=float(inst['pending_amount']))
                c9.number_format = '#,##0.00'
                c9.font = due_font if inst['pending_amount'] > 0 else normal_font
                c9.alignment = Alignment(horizontal="right")

                ws2.cell(row=cur_row2, column=10, value=inst['status']).alignment = Alignment(horizontal="center")

                for col in range(1, 11):
                    ws2.cell(row=cur_row2, column=col).border = THIN_BORDER

                cur_row2 += 1
                m_idx += 1

    if cur_row2 > start_row2:
        # Totals Row
        ws2.row_dimensions[cur_row2].height = 24
        ws2.merge_cells(f"A{cur_row2}:F{cur_row2}")
        ws2[f"A{cur_row2}"] = "TOTAL INSTALLMENTS SCHEDULE:"
        ws2[f"A{cur_row2}"].font = totals_bold
        ws2[f"A{cur_row2}"].alignment = Alignment(horizontal="right", vertical="center")

        t7 = ws2.cell(row=cur_row2, column=7, value=f"=SUM(G{start_row2}:G{cur_row2-1})")
        t7.number_format = '#,##0.00'
        t7.font = totals_bold
        t7.alignment = Alignment(horizontal="right", vertical="center")

        t8 = ws2.cell(row=cur_row2, column=8, value=f"=SUM(H{start_row2}:H{cur_row2-1})")
        t8.number_format = '#,##0.00'
        t8.font = totals_paid
        t8.alignment = Alignment(horizontal="right", vertical="center")

        t9 = ws2.cell(row=cur_row2, column=9, value=f"=SUM(I{start_row2}:I{cur_row2-1})")
        t9.number_format = '#,##0.00'
        t9.font = totals_due
        t9.alignment = Alignment(horizontal="right", vertical="center")

        for col in range(1, 11):
            cell = ws2.cell(row=cur_row2, column=col)
            cell.border = TOTALS_BORDER
            cell.fill = TOTALS_FILL

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # -------------------------------------------------------------
    # Sheet 3: Payment Receipts Ledger
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Receipts Ledger")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:K1")
    ws3["A1"] = f"{institute.get('name', 'SUTAWANE CLASSES').upper()} — COMPLETE PAYMENT RECEIPTS & AUDIT TRAIL"
    ws3["A1"].font = Font(name="Calibri", size=13, bold=True, color="1A2B4C")
    ws3["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[1].height = 26

    ws3.merge_cells("A2:K2")
    ws3["A2"] = f"Student: {student['full_name']} ({student['student_code']}) • Total Receipts: {len(payments_ledger)}"
    ws3["A2"].font = Font(name="Calibri", size=9.5, italic=True, color="64748B")
    ws3["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[2].height = 18

    headers3 = [
        "#", "Receipt Number", "Payment Date", "Academic Session", "Course / Batch",
        "Campus Branch", "Payment Mode", "Received By", "Milestone Allocation", "Amount (Rs.)", "Running Total (Rs.)"
    ]
    ws3.row_dimensions[4].height = 24
    for c_idx, h in enumerate(headers3, start=1):
        cell = ws3.cell(row=4, column=c_idx, value=h)
        cell.font = hdr_font
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="right" if "Rs." in h else "left", vertical="center")
        cell.border = THIN_BORDER

    cur_row3 = 5
    start_row3 = 5

    for p in payments_ledger:
        ws3.row_dimensions[cur_row3].height = 20
        ws3.cell(row=cur_row3, column=1, value=p['index']).alignment = Alignment(horizontal="center")
        ws3.cell(row=cur_row3, column=2, value=p['receipt_number']).font = bold_font
        
        p_dt = p['payment_date'].strftime('%d-%b-%Y') if p.get('payment_date') else '-'
        ws3.cell(row=cur_row3, column=3, value=p_dt).alignment = Alignment(horizontal="center")
        ws3.cell(row=cur_row3, column=4, value=p['session_name'])
        ws3.cell(row=cur_row3, column=5, value=p['standard'])
        ws3.cell(row=cur_row3, column=6, value=p['branch_name'])
        ws3.cell(row=cur_row3, column=7, value=p['payment_mode'])
        ws3.cell(row=cur_row3, column=8, value=p.get('received_by') or 'Staff')
        ws3.cell(row=cur_row3, column=9, value=p['allocations_summary'])

        c10 = ws3.cell(row=cur_row3, column=10, value=float(p['amount']))
        c10.number_format = '#,##0.00'
        c10.font = paid_font
        c10.alignment = Alignment(horizontal="right")

        c11 = ws3.cell(row=cur_row3, column=11, value=float(p['running_total']))
        c11.number_format = '#,##0.00'
        c11.font = totals_bold
        c11.alignment = Alignment(horizontal="right")

        for col in range(1, 12):
            ws3.cell(row=cur_row3, column=col).border = THIN_BORDER

        cur_row3 += 1

    if cur_row3 > start_row3:
        # Totals Row
        ws3.row_dimensions[cur_row3].height = 24
        ws3.merge_cells(f"A{cur_row3}:I{cur_row3}")
        ws3[f"A{cur_row3}"] = "GRAND TOTAL VERIFIED COLLECTIONS:"
        ws3[f"A{cur_row3}"].font = totals_bold
        ws3[f"A{cur_row3}"].alignment = Alignment(horizontal="right", vertical="center")

        t10 = ws3.cell(row=cur_row3, column=10, value=f"=SUM(J{start_row3}:J{cur_row3-1})")
        t10.number_format = '#,##0.00'
        t10.font = totals_paid
        t10.alignment = Alignment(horizontal="right", vertical="center")

        t11 = ws3.cell(row=cur_row3, column=11, value=float(summary['paid_fees']))
        t11.number_format = '#,##0.00'
        t11.font = totals_bold
        t11.alignment = Alignment(horizontal="right", vertical="center")

        for col in range(1, 12):
            cell = ws3.cell(row=cur_row3, column=col)
            cell.border = TOTALS_BORDER
            cell.fill = TOTALS_FILL

    for col in ws3.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws3.column_dimensions[col_letter].width = max(max_len + 4, 12)

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_data = excel_buffer.getvalue()
    excel_buffer.close()
    return excel_data
