"""
Export utilities for Center Expenses & Overheads Management.
Generates executive Landscape PDF and multi-sheet Excel (.xlsx) documents matching
the exact institutional snapshot design of Sutawane Classes software.
"""

import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from .export_theme import (
    COLOR_NAVY_PRIMARY,
    COLOR_NAVY_SECONDARY,
    COLOR_RED_PENDING,
    COLOR_BG_TOTALS,
    COLOR_BORDER,
    COLOR_SLATE_TEXT,
    format_currency_inr,
    build_executive_pdf_header,
    ExecutiveLandscapeCanvas,
    format_excel_worksheet,
    FONT_FAMILY,
    FONT_FAMILY_BOLD,
)


def export_expenses_to_excel(expenses, options=None):
    """
    Exports expense vouchers to an Openpyxl Excel (.xlsx) workbook.
    Features:
    - Multi-sheet data isolation:
      * Sheet 1: Consolidated Expense Vouchers
      * Sheet 2: Category Breakdown (Category, Count, Outflow Amount, Spend Share %)
      * Sheet 3+: Branch-specific sheets (when multiple branches present)
    - Row heights: Title=28, Meta=18, Header=24, Rows=20, Totals=24
    - Outflow amounts in bold Red (#DC2626)
    - Double border totals row with #F1F5F9 fill
    """
    options = options or {}
    branch_label = options.get('branch_name') or 'All Branches'
    session_label = options.get('session_name') or 'Current Session'
    now_str = datetime.now().strftime("%d-%b-%Y %I:%M %p")

    wb = Workbook()

    bold_font = Font(name="Calibri", size=10, bold=True)
    normal_font = Font(name="Calibri", size=10)
    outflow_font = Font(name="Calibri", size=10, bold=True, color="DC2626")
    totals_bold = Font(name="Calibri", size=11, bold=True, color="1A2B4C")
    totals_red = Font(name="Calibri", size=11, bold=True, color="DC2626")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    totals_border = Border(
        top=Side(style='medium', color='1A2B4C'),
        bottom=Side(style='double', color='1A2B4C'),
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1')
    )
    totals_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    headers = [
        ("#", 6),
        ("Expense Date", 14),
        ("Expense Category", 20),
        ("Voucher Description", 28),
        ("Branch Location", 20),
        ("Vendor / Payee", 20),
        ("Payment Mode", 15),
        ("Reference / UTR", 18),
        ("Recorded By", 20),
        ("Outflow Amount (₹)", 18),
    ]

    def populate_expense_sheet(ws, records, sheet_title, sheet_scope):
        ws.title = sheet_title[:31]
        header_row = 4

        for col_idx, (h_text, _) in enumerate(headers, start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=h_text)
            cell.alignment = Alignment(horizontal="center" if col_idx in [1, 2, 7, 8] else ("left" if col_idx in [3, 4, 5, 6, 9] else "right"), vertical="center")

        curr_row = 5
        tot_outflow = 0.0

        for idx, exp in enumerate(records, start=1):
            amt = float(exp.get('amount', 0.0))
            tot_outflow += amt
            rec_by = exp.get('recorded_by') or exp.get('created_by_display') or 'Staff'

            row_vals = [
                idx,
                exp.get('expense_date_formatted') or exp.get('expense_date_iso') or '—',
                exp.get('category_name') or 'General',
                exp.get('description') or '—',
                exp.get('branch_name') or branch_label,
                exp.get('vendor') or '—',
                exp.get('payment_mode') or 'UPI',
                exp.get('reference_number') or '—',
                rec_by,
                amt,
            ]

            ws.row_dimensions[curr_row].height = 20
            for col_idx, val in enumerate(row_vals, start=1):
                cell = ws.cell(row=curr_row, column=col_idx, value=val)
                cell.border = thin_border

                if col_idx == 10:
                    cell.number_format = '₹#,##0'
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                    cell.font = outflow_font
                elif col_idx in [1, 2, 7, 8]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    cell.font = normal_font
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                    cell.font = bold_font if col_idx == 3 else normal_font

            curr_row += 1

        # Totals Row
        ws.row_dimensions[curr_row].height = 24
        ws.merge_cells(start_row=curr_row, start_column=1, end_row=curr_row, end_column=9)
        lbl = ws.cell(row=curr_row, column=1, value=f"TOTAL OUTFLOW ({len(records)} vouchers)")
        lbl.font = totals_bold
        lbl.alignment = Alignment(horizontal="right", vertical="center")

        amt_cell = ws.cell(row=curr_row, column=10, value=tot_outflow)
        amt_cell.font = totals_red
        amt_cell.number_format = '₹#,##0'
        amt_cell.alignment = Alignment(horizontal="right", vertical="center")

        for col in range(1, 11):
            c = ws.cell(row=curr_row, column=col)
            c.border = totals_border
            c.fill = totals_fill

        meta_str = f"SCOPE: {sheet_scope}  |  DATE: {now_str}  |  TOTAL VOUCHERS: {len(records)}"
        format_excel_worksheet(ws, "SUTAWANE CLASSES — EXPENSE & OVERHEADS VOUCHERS REGISTER", meta_str, header_row=4)

    # 1. Main Sheet: All Vouchers
    ws_main = wb.active
    populate_expense_sheet(ws_main, expenses, "All Expenses", f"{branch_label} • {session_label}")

    # 2. Category-wise Breakdown Sheet
    ws_cat = wb.create_sheet(title="Category Breakdown")
    ws_cat.views.sheetView[0].showGridLines = True
    cat_headers = [("#", 6), ("Expense Category", 24), ("Vouchers Count", 16), ("Total Spent (₹)", 20), ("Share of Spend (%)", 20)]
    for col_idx, (h_text, _) in enumerate(cat_headers, start=1):
        cell = ws_cat.cell(row=4, column=col_idx, value=h_text)
        cell.alignment = Alignment(horizontal="center" if col_idx in [1, 3] else ("left" if col_idx == 2 else "right"), vertical="center")

    cat_map = {}
    grand_tot = sum(float(e.get('amount', 0.0)) for e in expenses)
    for e in expenses:
        cat = e.get('category_name') or 'General'
        if cat not in cat_map:
            cat_map[cat] = {'count': 0, 'amount': 0.0}
        cat_map[cat]['count'] += 1
        cat_map[cat]['amount'] += float(e.get('amount', 0.0))

    cat_row = 5
    for idx, (cat_name, c_data) in enumerate(sorted(cat_map.items(), key=lambda x: x[1]['amount'], reverse=True), start=1):
        ws_cat.row_dimensions[cat_row].height = 20
        share_pct = round((c_data['amount'] / grand_tot * 100), 1) if grand_tot > 0 else 0.0
        row_items = [idx, cat_name, c_data['count'], c_data['amount'], f"{share_pct}%"]
        for col_idx, val in enumerate(row_items, start=1):
            c = ws_cat.cell(row=cat_row, column=col_idx, value=val)
            c.border = thin_border
            if col_idx == 4:
                c.number_format = '₹#,##0'
                c.alignment = Alignment(horizontal="right", vertical="center")
                c.font = outflow_font
            elif col_idx in [1, 3, 5]:
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.font = bold_font if col_idx == 3 else normal_font
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")
                c.font = bold_font
        cat_row += 1

    # Totals Row for Category Sheet
    ws_cat.row_dimensions[cat_row].height = 24
    ws_cat.merge_cells(start_row=cat_row, start_column=1, end_row=cat_row, end_column=3)
    lbl_c = ws_cat.cell(row=cat_row, column=1, value="TOTAL CENTER OVERHEADS")
    lbl_c.font = totals_bold
    lbl_c.alignment = Alignment(horizontal="right", vertical="center")

    c_amt = ws_cat.cell(row=cat_row, column=4, value=grand_tot)
    c_amt.font = totals_red
    c_amt.number_format = '₹#,##0'
    c_amt.alignment = Alignment(horizontal="right", vertical="center")

    c_share = ws_cat.cell(row=cat_row, column=5, value="100.0%")
    c_share.font = totals_bold
    c_share.alignment = Alignment(horizontal="center", vertical="center")

    for col in range(1, 6):
        c = ws_cat.cell(row=cat_row, column=col)
        c.border = totals_border
        c.fill = totals_fill

    format_excel_worksheet(ws_cat, "SUTAWANE CLASSES — EXPENSE CATEGORY BREAKDOWN", f"SCOPE: {branch_label} • {session_label}  |  DATE: {now_str}", header_row=4)

    # 3. Dedicated Branch Sheets if multiple branches
    branches_present = sorted(list({e.get('branch_name') for e in expenses if e.get('branch_name') and e.get('branch_name') != '—'}))
    if len(branches_present) > 1:
        for b_name in branches_present:
            b_records = [e for e in expenses if e.get('branch_name') == b_name]
            if b_records:
                ws_b = wb.create_sheet(title=b_name[:31])
                populate_expense_sheet(ws_b, b_records, b_name, f"{b_name} • {session_label}")

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()


def export_expenses_to_pdf(expenses, options=None):
    """
    Exports expense vouchers to an executive Landscape A4 PDF document matching the snapshot.
    Features:
    - Official Header with Sutawane Logo, Title, Subtitle, Tagline
    - Institutional Scope Box on right
    - Amber and Navy Accent Bar
    - Outflow Amounts in bold Red (#DC2626)
    - Totals Row with matching colors and formatting
    - Two-pass Running Page Footer (Page X of Y)
    """
    options = options or {}
    institute_name = options.get('institute_name') or 'SUTAWANE CLASSES'
    branch_label = options.get('branch_name') or 'All Branches'
    session_label = options.get('session_name') or 'Current Session'
    gen_time = datetime.now().strftime('%d %b %Y, %I:%M %p')

    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=landscape(A4),
        leftMargin=18,
        rightMargin=18,
        topMargin=18,
        bottomMargin=26,
    )

    styles = getSampleStyleSheet()
    cell_normal = ParagraphStyle('CellNorm', parent=styles['Normal'], fontName=FONT_FAMILY, fontSize=8, leading=10, textColor=colors.HexColor(COLOR_SLATE_TEXT))
    cell_bold = ParagraphStyle('CellBld', parent=styles['Normal'], fontName=FONT_FAMILY_BOLD, fontSize=8, leading=10, textColor=colors.HexColor(COLOR_SLATE_TEXT))
    cell_red = ParagraphStyle('CellRed', parent=styles['Normal'], fontName=FONT_FAMILY_BOLD, fontSize=8, leading=10, textColor=colors.HexColor(COLOR_RED_PENDING), alignment=2)
    cell_center = ParagraphStyle('CellCtr', parent=styles['Normal'], fontName=FONT_FAMILY, fontSize=8, leading=10, alignment=1)
    cell_mode = ParagraphStyle('CellMode', parent=styles['Normal'], fontName=FONT_FAMILY_BOLD, fontSize=7.5, leading=9, textColor=colors.HexColor('#2563EB'), alignment=1)

    story = []

    # 1. Executive Snapshot Header
    meta_dict = {
        'institute_name': institute_name,
        'branch_name': branch_label,
        'session_name': session_label,
        'branch_logo': options.get('branch_logo'),
        'date_str': options.get('date_str') or gen_time,
        'records_count': len(expenses),
    }
    story.extend(build_executive_pdf_header(
        "EXPENSE & OVERHEADS VOUCHERS REPORT",
        meta=meta_dict,
        report_badge_label="OUTFLOWS:",
        report_badge_value=format_currency_inr(sum(float(e.get('amount', 0.0)) for e in expenses))
    ))

    # 2. Table Data
    table_data = [
        [
            "#",
            "Date",
            "Category",
            "Description",
            "Branch",
            "Vendor / Payee",
            "Payment Mode",
            "Ref / Bill No",
            "Recorded By",
            "Amount (Rs.)"
        ]
    ]

    total_amount = 0.0
    for idx, exp in enumerate(expenses, 1):
        amt = float(exp.get('amount', 0.0))
        total_amount += amt

        exp_date = exp.get('expense_date_formatted') or exp.get('expense_date_iso') or '—'
        cat_name = exp.get('category_name') or 'General'
        desc = exp.get('description') or '—'
        branch_name = exp.get('branch_name') or branch_label
        vendor = exp.get('vendor') or '—'
        mode = exp.get('payment_mode') or 'UPI'
        ref_no = exp.get('reference_number') or '—'
        rec_by = exp.get('recorded_by') or exp.get('created_by_display') or 'Staff'

        table_data.append([
            Paragraph(str(idx), cell_center),
            Paragraph(exp_date, cell_center),
            Paragraph(f"<b>{cat_name}</b>", cell_bold),
            Paragraph(desc, cell_normal),
            Paragraph(branch_name, cell_normal),
            Paragraph(vendor, cell_normal),
            Paragraph(mode, cell_mode),
            Paragraph(ref_no, cell_center),
            Paragraph(rec_by, cell_normal),
            Paragraph(format_currency_inr(amt), cell_red),
        ])

    # 3. Totals Row
    table_data.append([
        Paragraph(f"<b>TOTAL OUTFLOW ({len(expenses)} vouchers)</b>", cell_bold),
        "", "", "", "", "", "", "", "",
        Paragraph(f"<b>{format_currency_inr(total_amount)}</b>", cell_red),
    ])

    col_widths = [22, 62, 88, 138, 84, 84, 60, 76, 96, 94]

    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(COLOR_NAVY_SECONDARY)),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), FONT_FAMILY_BOLD),
        ('FONTSIZE', (0, 0), (-1, 0), 8),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 5.5),
        ('TOPPADDING', (0, 0), (-1, 0), 5.5),
        ('ALIGN', (0, 0), (0, -1), 'CENTER'),
        ('ALIGN', (1, 0), (1, -1), 'CENTER'),
        ('ALIGN', (6, 0), (7, -1), 'CENTER'),
        ('ALIGN', (9, 0), (9, -1), 'RIGHT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor(COLOR_BORDER)),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#F8FAFC')]),
        ('TOPPADDING', (0, 1), (-1, -1), 4.5),
        ('BOTTOMPADDING', (0, 1), (-1, -1), 4.5),
        ('SPAN', (0, -1), (8, -1)),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor(COLOR_BG_TOTALS)),
        ('LINEABOVE', (0, -1), (-1, -1), 1.25, colors.HexColor(COLOR_NAVY_PRIMARY)),
        ('LINEBELOW', (0, -1), (-1, -1), 1.25, colors.HexColor(COLOR_NAVY_PRIMARY)),
    ]))

    story.append(t)
    doc.build(story, canvasmaker=ExecutiveLandscapeCanvas)
    out.seek(0)
    return out.getvalue()
