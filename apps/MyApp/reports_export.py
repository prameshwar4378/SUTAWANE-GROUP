"""
Export utilities for Comprehensive Financial Reports, Day Book & Audit Statements.
Generates executive Landscape A4 PDF and multi-sheet Excel (.xlsx) documents matching
the exact institutional snapshot design of Sutawane Classes software.
"""

import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from .export_theme import (
    COLOR_NAVY_PRIMARY,
    COLOR_NAVY_SECONDARY,
    COLOR_AMBER_ACCENT,
    COLOR_GREEN_PAID,
    COLOR_RED_PENDING,
    COLOR_BG_CARD,
    COLOR_BG_TOTALS,
    COLOR_BORDER,
    COLOR_SLATE_TEXT,
    COLOR_MUTED_TEXT,
    format_currency_inr,
    format_number_inr,
    build_executive_pdf_header,
    ExecutiveLandscapeCanvas,
    format_excel_worksheet,
)


# Shared Openpyxl Border & Fill constants
THIN_BORDER = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1'),
)
TOTALS_BORDER = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='medium', color='1A2B4C'),
    bottom=Side(style='double', color='1A2B4C'),
)
HEADER_FILL = PatternFill(start_color="28375C", end_color="28375C", fill_type="solid")
TOTALS_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")


# =============================================================================
# 1. Session Audit Excel Export (Multi-Sheet Workbook)
# =============================================================================

def export_session_audit_excel(audit_data, options=None):
    """
    Exports comprehensive session audit to an Excel (.xlsx) workbook with:
    - Sheet 1: Branch Comparison (isolated branch KPIs)
    - Sheet 2: Course Breakdown (enrollment & revenue per course)
    - Sheet 3: Monthly Collection Trend
    - Row heights: Title=28, Meta=18, Header=24, Rows=20, Totals=24
    - Auto column width padding (+4 chars)
    - Green for collections, Red for pending dues & expenses
    """
    options = options or {}
    session_name = audit_data.get('session', {}).get('name') or options.get('session_name') or 'Current Session'
    branch_label = options.get('branch_name') or 'All Branches'
    now_str = datetime.now().strftime("%d-%b-%Y %I:%M %p")

    wb = Workbook()

    bold_font = Font(name="Calibri", size=10, bold=True)
    normal_font = Font(name="Calibri", size=10)
    paid_font = Font(name="Calibri", size=10, bold=True, color="16A34A")
    due_font = Font(name="Calibri", size=10, bold=True, color="DC2626")
    totals_bold = Font(name="Calibri", size=11, bold=True, color="1A2B4C")
    totals_paid = Font(name="Calibri", size=11, bold=True, color="16A34A")
    totals_due = Font(name="Calibri", size=11, bold=True, color="DC2626")

    # -----------------------------------------------------------------
    # Sheet 1: Branch Comparison
    # -----------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Branch Comparison"
    ws1.views.sheetView[0].showGridLines = True

    headers1 = [
        '#', 'Campus Branch', 'Code', 'Active Students', 'Agreed Fees (Rs.)',
        'Collections (Rs.)', 'Pending Dues (Rs.)', 'Center Expenses (Rs.)',
        'Net Surplus (Rs.)', 'Collection Rate (%)'
    ]

    max_col_letter1 = get_column_letter(len(headers1))
    ws1.merge_cells(f"A1:{max_col_letter1}1")
    ws1["A1"] = f"SUTAWANE CLASSES — ACADEMIC SESSION AUDIT: {session_name.upper()}"
    ws1["A1"].font = Font(name="Calibri", size=13, bold=True, color="1A2B4C")
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 28

    ws1.merge_cells(f"A2:{max_col_letter1}2")
    ws1["A2"] = f"Scope: {branch_label} | Generated on: {now_str}"
    ws1["A2"].font = Font(name="Calibri", size=9.5, italic=True, color="64748B")
    ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 18

    ws1.row_dimensions[3].height = 8

    ws1.row_dimensions[4].height = 24
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    for col_idx, h in enumerate(headers1, 1):
        cell = ws1.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        align_h = 'center' if col_idx in [1, 3, 4, 10] else ('right' if col_idx >= 5 else 'left')
        cell.alignment = Alignment(horizontal=align_h, vertical="center")

    branch_list = audit_data.get('branchBreakdown', [])
    tot_students = 0
    tot_fees = 0.0
    tot_col = 0.0
    tot_dues = 0.0
    tot_exp = 0.0

    for idx, b in enumerate(branch_list, 1):
        st_cnt = int(b.get('studentsCount', 0))
        fees = float(b.get('totalFees', 0.0))
        col = float(b.get('collections', 0.0))
        dues = float(b.get('pendingDues', 0.0))
        exp = float(b.get('expenses', 0.0))
        surplus = float(b.get('netSurplus', 0.0))
        rate = b.get('collectionRate', 0.0)

        tot_students += st_cnt
        tot_fees += fees
        tot_col += col
        tot_dues += dues
        tot_exp += exp

        row_num = 4 + idx
        ws1.row_dimensions[row_num].height = 20
        row_vals = [
            idx, b.get('branchName', '—'), b.get('branchCode', '—'),
            st_cnt, fees, col, dues, exp, surplus, f"{rate}%"
        ]

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws1.cell(row=row_num, column=col_idx, value=val)
            cell.border = THIN_BORDER

            if col_idx in [5, 6, 7, 8, 9]:
                cell.number_format = '₹#,##0'
                cell.alignment = Alignment(horizontal='right', vertical="center")
                if col_idx == 6:
                    cell.font = paid_font
                elif col_idx in [7, 8]:
                    cell.font = due_font if val > 0 else normal_font
                else:
                    cell.font = bold_font
            elif col_idx in [1, 3, 4, 10]:
                cell.font = normal_font
                cell.alignment = Alignment(horizontal='center', vertical="center")
            else:
                cell.font = bold_font if col_idx == 2 else normal_font
                cell.alignment = Alignment(horizontal='left', vertical="center")

    # Summary Row Sheet 1
    tot_surplus = tot_col - tot_exp
    tot_rate = round((tot_col / tot_fees * 100), 1) if tot_fees > 0 else 0.0
    tot_row_num = 4 + len(branch_list) + 1
    ws1.row_dimensions[tot_row_num].height = 24

    sum_vals1 = [
        'TOTAL', 'Consolidated Summary', '', tot_students, tot_fees,
        tot_col, tot_dues, tot_exp, tot_surplus, f"{tot_rate}%"
    ]
    for col_idx, val in enumerate(sum_vals1, 1):
        cell = ws1.cell(row=tot_row_num, column=col_idx, value=val)
        cell.border = TOTALS_BORDER
        cell.fill = TOTALS_FILL

        if col_idx in [5, 6, 7, 8, 9]:
            cell.number_format = '₹#,##0'
            cell.alignment = Alignment(horizontal='right', vertical="center")
            if col_idx == 6:
                cell.font = totals_paid
            elif col_idx in [7, 8]:
                cell.font = totals_due
            else:
                cell.font = totals_bold
        elif col_idx in [1, 2]:
            cell.font = totals_bold
            cell.alignment = Alignment(horizontal='left', vertical="center")
        else:
            cell.font = totals_bold
            cell.alignment = Alignment(horizontal='center', vertical="center")

    # Column widths
    for col in ws1.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row in (1, 2, 3):
                continue
            v_str = str(cell.value or '')
            if len(v_str) > max_len:
                max_len = len(v_str)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # -----------------------------------------------------------------
    # Sheet 2: Course Breakdown
    # -----------------------------------------------------------------
    ws2 = wb.create_sheet(title="Course Breakdown")
    ws2.views.sheetView[0].showGridLines = True

    headers2 = ['#', 'Course / Standard', 'Enrolled Students', 'Total Committed Fees (Rs.)', 'Collections (Rs.)', 'Pending Fees (Rs.)', 'Collection Rate (%)']
    max_col_letter2 = get_column_letter(len(headers2))

    ws2.merge_cells(f"A1:{max_col_letter2}1")
    ws2["A1"] = f"COURSE-WISE ENROLLMENT & REVENUE RECONCILIATION — {session_name.upper()}"
    ws2["A1"].font = Font(name="Calibri", size=13, bold=True, color="1A2B4C")
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 28

    ws2.merge_cells(f"A2:{max_col_letter2}2")
    ws2["A2"] = f"Scope: {branch_label} | Generated on: {now_str}"
    ws2["A2"].font = Font(name="Calibri", size=9.5, italic=True, color="64748B")
    ws2["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[2].height = 18

    ws2.row_dimensions[3].height = 8

    ws2.row_dimensions[4].height = 24
    for col_idx, h in enumerate(headers2, 1):
        cell = ws2.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        align_h = 'center' if col_idx in [1, 3, 7] else ('right' if col_idx in [4, 5, 6] else 'left')
        cell.alignment = Alignment(horizontal=align_h, vertical="center")

    course_list = audit_data.get('courseBreakdown', [])
    c_tot_st = 0
    c_tot_fees = 0.0
    c_tot_col = 0.0
    c_tot_dues = 0.0

    for idx, c_item in enumerate(course_list, 1):
        st_cnt = int(c_item.get('studentsCount', 0))
        fees = float(c_item.get('totalFees', 0.0))
        col = float(c_item.get('collections', 0.0))
        dues = float(c_item.get('pendingDues', 0.0))
        rate = c_item.get('collectionRate', 0.0)

        c_tot_st += st_cnt
        c_tot_fees += fees
        c_tot_col += col
        c_tot_dues += dues

        row_num = 4 + idx
        ws2.row_dimensions[row_num].height = 20
        row_vals = [idx, c_item.get('standard', '—'), st_cnt, fees, col, dues, f"{rate}%"]

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws2.cell(row=row_num, column=col_idx, value=val)
            cell.border = THIN_BORDER

            if col_idx in [4, 5, 6]:
                cell.number_format = '₹#,##0'
                cell.alignment = Alignment(horizontal='right', vertical="center")
                if col_idx == 5:
                    cell.font = paid_font
                elif col_idx == 6:
                    cell.font = due_font if val > 0 else normal_font
                else:
                    cell.font = bold_font
            elif col_idx in [1, 3, 7]:
                cell.font = normal_font
                cell.alignment = Alignment(horizontal='center', vertical="center")
            else:
                cell.font = bold_font
                cell.alignment = Alignment(horizontal='left', vertical="center")

    # Summary Row Sheet 2
    c_tot_rate = round((c_tot_col / c_tot_fees * 100), 1) if c_tot_fees > 0 else 0.0
    tot_row_num2 = 4 + len(course_list) + 1
    ws2.row_dimensions[tot_row_num2].height = 24

    sum_vals2 = ['TOTAL', 'Consolidated Courses', c_tot_st, c_tot_fees, c_tot_col, c_tot_dues, f"{c_tot_rate}%"]
    for col_idx, val in enumerate(sum_vals2, 1):
        cell = ws2.cell(row=tot_row_num2, column=col_idx, value=val)
        cell.border = TOTALS_BORDER
        cell.fill = TOTALS_FILL

        if col_idx in [4, 5, 6]:
            cell.number_format = '₹#,##0'
            cell.alignment = Alignment(horizontal='right', vertical="center")
            if col_idx == 5:
                cell.font = totals_paid
            elif col_idx == 6:
                cell.font = totals_due
            else:
                cell.font = totals_bold
        elif col_idx in [1, 2]:
            cell.font = totals_bold
            cell.alignment = Alignment(horizontal='left', vertical="center")
        else:
            cell.font = totals_bold
            cell.alignment = Alignment(horizontal='center', vertical="center")

    for col in ws2.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row in (1, 2, 3):
                continue
            v_str = str(cell.value or '')
            if len(v_str) > max_len:
                max_len = len(v_str)
        ws2.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # -----------------------------------------------------------------
    # Sheet 3: Monthly Trend
    # -----------------------------------------------------------------
    ws3 = wb.create_sheet(title="Monthly Trend")
    ws3.views.sheetView[0].showGridLines = True

    headers3 = ['#', 'Month', 'Collections (Rs.)', 'Disbursements (Rs.)', 'Net Monthly Surplus (Rs.)']
    max_col_letter3 = get_column_letter(len(headers3))

    ws3.merge_cells(f"A1:{max_col_letter3}1")
    ws3["A1"] = f"MONTHLY CASH INFLOW VS OUTFLOW TREND — {session_name.upper()}"
    ws3["A1"].font = Font(name="Calibri", size=13, bold=True, color="1A2B4C")
    ws3["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[1].height = 28

    ws3.merge_cells(f"A2:{max_col_letter3}2")
    ws3["A2"] = f"Scope: {branch_label} | Generated on: {now_str}"
    ws3["A2"].font = Font(name="Calibri", size=9.5, italic=True, color="64748B")
    ws3["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[2].height = 18

    ws3.row_dimensions[3].height = 8

    ws3.row_dimensions[4].height = 24
    for col_idx, h in enumerate(headers3, 1):
        cell = ws3.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        align_h = 'center' if col_idx in [1, 2] else 'right'
        cell.alignment = Alignment(horizontal=align_h, vertical="center")

    monthly_list = audit_data.get('monthlyBreakdown', [])
    m_tot_col = 0.0
    m_tot_exp = 0.0

    for idx, m_item in enumerate(monthly_list, 1):
        m_col = float(m_item.get('collections', 0.0))
        m_exp = float(m_item.get('expenses', 0.0))
        m_net = m_col - m_exp
        m_tot_col += m_col
        m_tot_exp += m_exp

        row_num = 4 + idx
        ws3.row_dimensions[row_num].height = 20
        row_vals = [idx, m_item.get('month', '—'), m_col, m_exp, m_net]

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws3.cell(row=row_num, column=col_idx, value=val)
            cell.border = THIN_BORDER

            if col_idx in [3, 4, 5]:
                cell.number_format = '₹#,##0'
                cell.alignment = Alignment(horizontal='right', vertical="center")
                if col_idx == 3:
                    cell.font = paid_font
                elif col_idx == 4:
                    cell.font = due_font
                else:
                    cell.font = paid_font if val >= 0 else due_font
            elif col_idx in [1, 2]:
                cell.font = bold_font if col_idx == 2 else normal_font
                cell.alignment = Alignment(horizontal='center', vertical="center")

    # Summary Row Sheet 3
    tot_row_num3 = 4 + len(monthly_list) + 1
    ws3.row_dimensions[tot_row_num3].height = 24
    sum_vals3 = ['TOTAL', 'Full Session Trend', m_tot_col, m_tot_exp, (m_tot_col - m_tot_exp)]

    for col_idx, val in enumerate(sum_vals3, 1):
        cell = ws3.cell(row=tot_row_num3, column=col_idx, value=val)
        cell.border = TOTALS_BORDER
        cell.fill = TOTALS_FILL

        if col_idx in [3, 4, 5]:
            cell.number_format = '₹#,##0'
            cell.alignment = Alignment(horizontal='right', vertical="center")
            if col_idx == 3:
                cell.font = totals_paid
            elif col_idx == 4:
                cell.font = totals_due
            else:
                cell.font = totals_bold
        elif col_idx in [1, 2]:
            cell.font = totals_bold
            cell.alignment = Alignment(horizontal='left' if col_idx == 2 else 'center', vertical="center")

    for col in ws3.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row in (1, 2, 3):
                continue
            v_str = str(cell.value or '')
            if len(v_str) > max_len:
                max_len = len(v_str)
        ws3.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


# =============================================================================
# 2. Session Audit PDF Export (Landscape A4)
# =============================================================================

def export_session_audit_pdf(audit_data, options=None):
    """
    Generates high-grade Landscape A4 Academic Session Audit PDF document matching the snapshot layout:
    - Title banner matching institutional design
    - Branch-by-branch financial performance comparison
    - Course-wise enrollment and fees distribution
    - Green for collections, Red for pending dues & expenses
    """
    options = options or {}
    session_name = audit_data.get('session', {}).get('name') or options.get('session_name') or 'Current Session'
    branch_label = options.get('branch_name') or 'All Branches'

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=20,
        rightMargin=20,
        topMargin=20,
        bottomMargin=30,
    )

    styles = getSampleStyleSheet()

    hdr_style = ParagraphStyle(
        'AudHdr',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.white,
        alignment=1,
    )
    hdr_right_style = ParagraphStyle(
        'AudHdrR',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.white,
        alignment=2,
    )
    section_title = ParagraphStyle(
        'AudSecTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9.5,
        textColor=colors.HexColor(COLOR_NAVY_PRIMARY),
        leading=13,
    )
    cell_normal = ParagraphStyle(
        'AudCellNorm',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7,
        leading=9,
        textColor=colors.HexColor(COLOR_SLATE_TEXT),
    )
    cell_bold = ParagraphStyle(
        'AudCellBld',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7,
        leading=9,
        textColor=colors.HexColor(COLOR_SLATE_TEXT),
    )
    cell_center = ParagraphStyle(
        'AudCellCenter',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7,
        leading=9,
        textColor=colors.HexColor(COLOR_SLATE_TEXT),
        alignment=1,
    )
    cell_right = ParagraphStyle(
        'AudCellRight',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7,
        leading=9,
        textColor=colors.HexColor(COLOR_SLATE_TEXT),
        alignment=2,
    )
    cell_green = ParagraphStyle(
        'AudCellGreen',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7,
        leading=9,
        textColor=colors.HexColor(COLOR_GREEN_PAID),
        alignment=2,
    )
    cell_red = ParagraphStyle(
        'AudCellRed',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7,
        leading=9,
        textColor=colors.HexColor(COLOR_RED_PENDING),
        alignment=2,
    )

    story = []

    # 1. Executive Snapshot Header
    branch_list = audit_data.get('branchBreakdown', [])
    header_meta = {
        'institute_name': 'SUTAWANE CLASSES VAIJAPUR',
        'branch_name': branch_label,
        'session_name': session_name,
        'records_count': len(branch_list),
    }
    story.extend(build_executive_pdf_header(
        doc_title="ACADEMIC SESSION AUDIT & ENROLLMENT REPORT",
        meta=header_meta,
        report_badge_label="STATUS:",
        report_badge_value="Audit Reconciled",
    ))

    # 2. Branch Performance Table
    story.append(Paragraph(f"<b>Part A: Campus Branch Financial Performance ({len(branch_list)} Centers)</b>", section_title))
    story.append(Spacer(1, 4))

    b_headers = [
        Paragraph("#", hdr_style),
        Paragraph("Campus Branch", hdr_style),
        Paragraph("Code", hdr_style),
        Paragraph("Students", hdr_style),
        Paragraph("Committed Fees", hdr_right_style),
        Paragraph("Collections", hdr_right_style),
        Paragraph("Pending Dues", hdr_right_style),
        Paragraph("Center Expenses", hdr_right_style),
        Paragraph("Net Surplus", hdr_right_style),
        Paragraph("Rate", hdr_style),
    ]
    b_rows = [b_headers]

    tot_st = 0
    tot_fees = 0.0
    tot_col = 0.0
    tot_dues = 0.0
    tot_exp = 0.0

    for idx, b in enumerate(branch_list, 1):
        st_cnt = int(b.get('studentsCount', 0))
        fees = float(b.get('totalFees', 0.0))
        col = float(b.get('collections', 0.0))
        dues = float(b.get('pendingDues', 0.0))
        exp = float(b.get('expenses', 0.0))
        surplus = float(b.get('netSurplus', 0.0))
        rate = b.get('collectionRate', 0.0)

        tot_st += st_cnt
        tot_fees += fees
        tot_col += col
        tot_dues += dues
        tot_exp += exp

        b_rows.append([
            Paragraph(str(idx), cell_center),
            Paragraph(str(b.get('branchName', '—')), cell_bold),
            Paragraph(str(b.get('branchCode', '—')), cell_center),
            Paragraph(str(st_cnt), cell_center),
            Paragraph(format_currency_inr(fees), cell_right),
            Paragraph(format_currency_inr(col), cell_green),
            Paragraph(format_currency_inr(dues), cell_red if dues > 0 else cell_right),
            Paragraph(format_currency_inr(exp), cell_red if exp > 0 else cell_right),
            Paragraph(format_currency_inr(surplus), cell_green if surplus >= 0 else cell_red),
            Paragraph(f"{rate}%", cell_center),
        ])

    tot_surplus = tot_col - tot_exp
    tot_rate = round((tot_col / tot_fees * 100), 1) if tot_fees > 0 else 0.0

    b_rows.append([
        Paragraph("<b>TOTAL</b>", cell_bold),
        Paragraph("<b>Consolidated</b>", cell_bold),
        Paragraph("", cell_normal),
        Paragraph(f"<b>{tot_st}</b>", cell_center),
        Paragraph(format_currency_inr(tot_fees), cell_bold),
        Paragraph(format_currency_inr(tot_col), cell_green),
        Paragraph(format_currency_inr(tot_dues), cell_red),
        Paragraph(format_currency_inr(tot_exp), cell_red),
        Paragraph(format_currency_inr(tot_surplus), cell_green if tot_surplus >= 0 else cell_red),
        Paragraph(f"<b>{tot_rate}%</b>", cell_center),
    ])

    # Col widths totaling 802 pt
    b_table = Table(b_rows, colWidths=[24, 148, 55, 55, 90, 90, 90, 90, 100, 60])
    b_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(COLOR_NAVY_SECONDARY)),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor(COLOR_BORDER)),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor(COLOR_BG_TOTALS)),
        ('LINEABOVE', (0, -1), (-1, -1), 1.2, colors.HexColor(COLOR_NAVY_PRIMARY)),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#F8FAFC')]),
    ]))
    story.append(b_table)
    story.append(Spacer(1, 12))

    # 3. Course Breakdown Table
    course_list = audit_data.get('courseBreakdown', [])
    story.append(Paragraph(f"<b>Part B: Course & Standard Wise Revenue Breakdown ({len(course_list)} Courses)</b>", section_title))
    story.append(Spacer(1, 4))

    c_headers = [
        Paragraph("#", hdr_style),
        Paragraph("Course / Standard", hdr_style),
        Paragraph("Enrolled Students", hdr_style),
        Paragraph("Total Committed Fees", hdr_right_style),
        Paragraph("Collections Realized", hdr_right_style),
        Paragraph("Pending Balances", hdr_right_style),
        Paragraph("Collection Efficiency", hdr_style),
    ]
    c_rows = [c_headers]

    c_tot_st = 0
    c_tot_fees = 0.0
    c_tot_col = 0.0
    c_tot_dues = 0.0

    for idx, c_item in enumerate(course_list, 1):
        st_cnt = int(c_item.get('studentsCount', 0))
        fees = float(c_item.get('totalFees', 0.0))
        col = float(c_item.get('collections', 0.0))
        dues = float(c_item.get('pendingDues', 0.0))
        rate = c_item.get('collectionRate', 0.0)

        c_tot_st += st_cnt
        c_tot_fees += fees
        c_tot_col += col
        c_tot_dues += dues

        c_rows.append([
            Paragraph(str(idx), cell_center),
            Paragraph(str(c_item.get('standard', '—')), cell_bold),
            Paragraph(str(st_cnt), cell_center),
            Paragraph(format_currency_inr(fees), cell_right),
            Paragraph(format_currency_inr(col), cell_green),
            Paragraph(format_currency_inr(dues), cell_red if dues > 0 else cell_right),
            Paragraph(f"{rate}%", cell_center),
        ])

    c_tot_rate = round((c_tot_col / c_tot_fees * 100), 1) if c_tot_fees > 0 else 0.0
    c_rows.append([
        Paragraph("<b>TOTAL</b>", cell_bold),
        Paragraph("<b>Consolidated Courses</b>", cell_bold),
        Paragraph(f"<b>{c_tot_st}</b>", cell_center),
        Paragraph(format_currency_inr(c_tot_fees), cell_bold),
        Paragraph(format_currency_inr(c_tot_col), cell_green),
        Paragraph(format_currency_inr(c_tot_dues), cell_red),
        Paragraph(f"<b>{c_tot_rate}%</b>", cell_center),
    ])

    c_table = Table(c_rows, colWidths=[24, 218, 100, 120, 120, 120, 100])
    c_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(COLOR_NAVY_SECONDARY)),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor(COLOR_BORDER)),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor(COLOR_BG_TOTALS)),
        ('LINEABOVE', (0, -1), (-1, -1), 1.2, colors.HexColor(COLOR_NAVY_PRIMARY)),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#F8FAFC')]),
    ]))
    story.append(c_table)

    doc.build(story, canvasmaker=ExecutiveLandscapeCanvas)
    buffer.seek(0)
    return buffer.getvalue()


# =============================================================================
# 3. Payment Modes Excel Export (Multi-Sheet)
# =============================================================================

def export_payment_modes_excel(modes_data, options=None):
    """
    Exports payment modes reconciliation to an Excel (.xlsx) file with:
    - Sheet 1: Payment Channels Summary
    - Sheet 2: Channel Insights & Audit
    - Row heights: Title=28, Meta=18, Header=24, Rows=20, Totals=24
    - Green font for Realized Amounts (#16A34A)
    - Auto column widths
    """
    options = options or {}
    branch_label = options.get('branch_name') or 'All Branches'
    session_label = options.get('session_name') or 'All Sessions'
    now_str = datetime.now().strftime("%d-%b-%Y %I:%M %p")

    wb = Workbook()

    bold_font = Font(name="Calibri", size=10, bold=True)
    normal_font = Font(name="Calibri", size=10)
    paid_font = Font(name="Calibri", size=10, bold=True, color="16A34A")
    totals_bold = Font(name="Calibri", size=11, bold=True, color="1A2B4C")
    totals_paid = Font(name="Calibri", size=11, bold=True, color="16A34A")

    # Sheet 1: Channels Summary
    ws1 = wb.active
    ws1.title = "Payment Channels"
    ws1.views.sheetView[0].showGridLines = True

    headers = ['#', 'Payment Channel / Mode', 'Total Realized (Rs.)', 'Receipts Count', 'Average Transaction (Rs.)', 'Share of Total (%)']
    max_col_letter = get_column_letter(len(headers))

    ws1.merge_cells(f"A1:{max_col_letter}1")
    ws1["A1"] = "SUTAWANE CLASSES — RECONCILED PAYMENT MODES REGISTER"
    ws1["A1"].font = Font(name="Calibri", size=13, bold=True, color="1A2B4C")
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 28

    ws1.merge_cells(f"A2:{max_col_letter}2")
    ws1["A2"] = f"Branch: {branch_label} | Session: {session_label} | Generated on: {now_str}"
    ws1["A2"].font = Font(name="Calibri", size=9.5, italic=True, color="64748B")
    ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 18

    ws1.row_dimensions[3].height = 8

    ws1.row_dimensions[4].height = 24
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    for col_idx, h in enumerate(headers, 1):
        cell = ws1.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        align_h = 'center' if col_idx in [1, 4, 6] else ('right' if col_idx in [3, 5] else 'left')
        cell.alignment = Alignment(horizontal=align_h, vertical="center")

    modes = modes_data.get('modes', [])
    tot_amt = 0.0
    tot_cnt = 0

    for idx, m in enumerate(modes, 1):
        amt = float(m.get('amount', 0.0))
        cnt = int(m.get('count', 0))
        avg = round(amt / cnt, 2) if cnt > 0 else 0.0
        pct = m.get('percentage', 0.0)

        tot_amt += amt
        tot_cnt += cnt

        row_num = 4 + idx
        ws1.row_dimensions[row_num].height = 20
        row_vals = [idx, m.get('label', '—'), amt, cnt, avg, f"{pct}%"]

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws1.cell(row=row_num, column=col_idx, value=val)
            cell.border = THIN_BORDER

            if col_idx in [3, 5]:
                cell.number_format = '₹#,##0'
                cell.alignment = Alignment(horizontal='right', vertical="center")
                cell.font = paid_font if col_idx == 3 else bold_font
            elif col_idx in [1, 4, 6]:
                cell.font = normal_font
                cell.alignment = Alignment(horizontal='center', vertical="center")
            else:
                cell.font = bold_font
                cell.alignment = Alignment(horizontal='left', vertical="center")

    # Summary Row
    overall_avg = round(tot_amt / tot_cnt, 2) if tot_cnt > 0 else 0.0
    tot_row_num = 4 + len(modes) + 1
    ws1.row_dimensions[tot_row_num].height = 24

    sum_vals = ['TOTAL', 'Consolidated Inflows', tot_amt, tot_cnt, overall_avg, '100%']
    for col_idx, val in enumerate(sum_vals, 1):
        cell = ws1.cell(row=tot_row_num, column=col_idx, value=val)
        cell.border = TOTALS_BORDER
        cell.fill = TOTALS_FILL

        if col_idx in [3, 5]:
            cell.number_format = '₹#,##0'
            cell.alignment = Alignment(horizontal='right', vertical="center")
            cell.font = totals_paid if col_idx == 3 else totals_bold
        elif col_idx in [1, 2]:
            cell.font = totals_bold
            cell.alignment = Alignment(horizontal='left', vertical="center")
        else:
            cell.font = totals_bold
            cell.alignment = Alignment(horizontal='center', vertical="center")

    for col in ws1.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row in (1, 2, 3):
                continue
            v_str = str(cell.value or '')
            if len(v_str) > max_len:
                max_len = len(v_str)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # Sheet 2: Channel Insights
    ws2 = wb.create_sheet(title="Channel Insights")
    ws2.views.sheetView[0].showGridLines = True
    ws2.merge_cells("A1:D1")
    ws2["A1"] = "PAYMENT METHOD DISTRIBUTION INSIGHTS"
    ws2["A1"].font = Font(name="Calibri", size=12, bold=True, color="1A2B4C")
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 26

    ws2.merge_cells("A2:D2")
    ws2["A2"] = f"Scope: {branch_label} • Reconciled on {now_str}"
    ws2["A2"].font = Font(name="Calibri", size=9, italic=True, color="64748B")
    ws2["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[2].height = 18

    ws2.row_dimensions[3].height = 8

    ins_headers = ['Metric Key', 'Observed Value', 'Benchmark / Target', 'Remarks']
    ws2.row_dimensions[4].height = 22
    for c_idx, h in enumerate(ins_headers, 1):
        c = ws2.cell(row=4, column=c_idx, value=h)
        c.font = header_font
        c.fill = HEADER_FILL
        c.border = THIN_BORDER
        c.alignment = Alignment(horizontal="center" if c_idx in [1, 3] else "left", vertical="center")

    digital_modes = sum(m.get('amount', 0.0) for m in modes if m.get('label', '').upper() in ('UPI', 'ONLINE', 'BANK TRANSFER', 'CARD', 'NET BANKING'))
    cash_mode = sum(m.get('amount', 0.0) for m in modes if m.get('label', '').upper() == 'CASH')
    digital_share = round((digital_modes / tot_amt * 100), 1) if tot_amt > 0 else 0.0

    insight_rows = [
        ['Total Realized Inflows', format_currency_inr(tot_amt), '100% Target', 'Reconciled to bank & cash accounts'],
        ['Total Payment Vouchers', str(tot_cnt), 'N/A', 'Official receipts issued'],
        ['Digital Collection Share (UPI/Bank)', f"{digital_share}%", '> 60% Optimal', f"Realized: {format_currency_inr(digital_modes)}"],
        ['Cash Collection Share', f"{round(100 - digital_share, 1)}%", '< 40% Optimal', f"Realized: {format_currency_inr(cash_mode)}"],
        ['Average Receipt Size', format_currency_inr(overall_avg), 'Ticket size indicator', 'Across all active channels'],
    ]

    for r_idx, r_data in enumerate(insight_rows, 5):
        ws2.row_dimensions[r_idx].height = 20
        for c_idx, val in enumerate(r_data, 1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=val)
            cell.border = THIN_BORDER
            cell.font = bold_font if c_idx in [1, 2] else normal_font
            cell.alignment = Alignment(horizontal="center" if c_idx == 3 else "left", vertical="center")

    for col in ws2.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(c.value or '')) for c in col if c.row > 3) or 15
        ws2.column_dimensions[col_letter].width = max_len + 4

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


# =============================================================================
# 4. Payment Modes PDF Export (Landscape A4 - Fixed from Portrait)
# =============================================================================

def export_payment_modes_pdf(modes_data, options=None):
    """
    Generates Landscape A4 Payment Modes Summary PDF report matching official snapshot layout:
    - Institutional header with logo, title, and scope box
    - Landscape table with Green realized amounts (#16A34A)
    - Reconciled totals row
    - Two-pass page numbering
    """
    options = options or {}
    branch_label = options.get('branch_name') or 'All Campuses'
    session_label = options.get('session_name') or 'Academic Session'

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=20,
        rightMargin=20,
        topMargin=20,
        bottomMargin=30,
    )

    styles = getSampleStyleSheet()

    hdr_style = ParagraphStyle(
        'ModHdr',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=1,
    )
    hdr_right_style = ParagraphStyle(
        'ModHdrR',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=2,
    )
    cell_normal = ParagraphStyle(
        'ModCellNorm',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor(COLOR_SLATE_TEXT),
    )
    cell_bold = ParagraphStyle(
        'ModCellBld',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor(COLOR_SLATE_TEXT),
    )
    cell_center = ParagraphStyle(
        'ModCellCenter',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor(COLOR_SLATE_TEXT),
        alignment=1,
    )
    cell_right = ParagraphStyle(
        'ModCellRight',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor(COLOR_SLATE_TEXT),
        alignment=2,
    )
    cell_green = ParagraphStyle(
        'ModCellGreen',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor(COLOR_GREEN_PAID),
        alignment=2,
    )

    story = []

    # 1. Executive Snapshot Header
    modes = modes_data.get('modes', [])
    tot_col = float(modes_data.get('totalCollected', 0.0))
    tot_cnt = int(modes_data.get('totalTransactions', 0))

    header_meta = {
        'institute_name': 'SUTAWANE CLASSES VAIJAPUR',
        'branch_name': branch_label,
        'session_name': session_label,
        'records_count': len(modes),
    }
    story.extend(build_executive_pdf_header(
        doc_title="PAYMENT CHANNELS & INFLOW RECONCILIATION",
        meta=header_meta,
        report_badge_label="REALIZED:",
        report_badge_value=format_currency_inr(tot_col),
    ))

    # 2. Table Headers
    tbl_headers = [
        Paragraph("#", hdr_style),
        Paragraph("Payment Channel / Mode", hdr_style),
        Paragraph("Realized Amount (Rs.)", hdr_right_style),
        Paragraph("Receipts Count", hdr_style),
        Paragraph("Average / Receipt (Rs.)", hdr_right_style),
        Paragraph("Share (%)", hdr_style),
    ]
    tbl_rows = [tbl_headers]

    for idx, m in enumerate(modes, 1):
        amt = float(m.get('amount', 0.0))
        cnt = int(m.get('count', 0))
        avg = round(amt / cnt, 2) if cnt > 0 else 0.0
        pct = m.get('percentage', 0.0)

        tbl_rows.append([
            Paragraph(str(idx), cell_center),
            Paragraph(str(m.get('label', '—')), cell_bold),
            Paragraph(format_currency_inr(amt), cell_green),
            Paragraph(str(cnt), cell_center),
            Paragraph(format_currency_inr(avg), cell_right),
            Paragraph(f"{pct}%", cell_center),
        ])

    overall_avg = round(tot_col / tot_cnt, 2) if tot_cnt > 0 else 0.0
    tbl_rows.append([
        Paragraph("<b>TOTAL</b>", cell_bold),
        Paragraph("<b>Consolidated Inflows</b>", cell_bold),
        Paragraph(format_currency_inr(tot_col), cell_green),
        Paragraph(f"<b>{tot_cnt}</b>", cell_center),
        Paragraph(format_currency_inr(overall_avg), cell_bold),
        Paragraph("<b>100%</b>", cell_center),
    ])

    # Col widths totaling 802 pt (Landscape A4)
    # [30, 222, 160, 110, 150, 130] = 802
    mode_tbl = Table(tbl_rows, colWidths=[30, 222, 160, 110, 150, 130])
    mode_tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(COLOR_NAVY_SECONDARY)),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor(COLOR_BORDER)),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor(COLOR_BG_TOTALS)),
        ('LINEABOVE', (0, -1), (-1, -1), 1.2, colors.HexColor(COLOR_NAVY_PRIMARY)),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#F8FAFC')]),
    ]))
    story.append(mode_tbl)

    doc.build(story, canvasmaker=ExecutiveLandscapeCanvas)
    buffer.seek(0)
    return buffer.getvalue()


# =============================================================================
# 5. Expense Categories Excel Export (Multi-Sheet)
# =============================================================================

def export_expense_categories_excel(categories_data, options=None):
    """
    Exports expense categories taxonomy and distribution to an Excel (.xlsx) file with:
    - Sheet 1: Expense Categories Breakdown
    - Sheet 2: Budget Distribution Insights
    - Row heights: Title=28, Meta=18, Header=24, Rows=20, Totals=24
    - Red font for Disbursed Amounts (#DC2626)
    - Auto column widths
    """
    options = options or {}
    branch_label = options.get('branch_name') or 'All Branches'
    now_str = datetime.now().strftime("%d-%b-%Y %I:%M %p")

    wb = Workbook()

    bold_font = Font(name="Calibri", size=10, bold=True)
    normal_font = Font(name="Calibri", size=10)
    due_font = Font(name="Calibri", size=10, bold=True, color="DC2626")
    totals_bold = Font(name="Calibri", size=11, bold=True, color="1A2B4C")
    totals_red = Font(name="Calibri", size=11, bold=True, color="DC2626")

    # Sheet 1: Categories Summary
    ws1 = wb.active
    ws1.title = "Expense Categories"
    ws1.views.sheetView[0].showGridLines = True

    headers = ['#', 'Expense Category', 'Total Disbursed (Rs.)', 'Vouchers Count', 'Average / Voucher (Rs.)', 'Budget Share (%)']
    max_col_letter = get_column_letter(len(headers))

    ws1.merge_cells(f"A1:{max_col_letter}1")
    ws1["A1"] = "SUTAWANE CLASSES — OPERATIONAL EXPENSES BY CATEGORY"
    ws1["A1"].font = Font(name="Calibri", size=13, bold=True, color="1A2B4C")
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 28

    ws1.merge_cells(f"A2:{max_col_letter}2")
    ws1["A2"] = f"Branch: {branch_label} | Generated on: {now_str}"
    ws1["A2"].font = Font(name="Calibri", size=9.5, italic=True, color="64748B")
    ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 18

    ws1.row_dimensions[3].height = 8

    ws1.row_dimensions[4].height = 24
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    for col_idx, h in enumerate(headers, 1):
        cell = ws1.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        align_h = 'center' if col_idx in [1, 4, 6] else ('right' if col_idx in [3, 5] else 'left')
        cell.alignment = Alignment(horizontal=align_h, vertical="center")

    cats = categories_data.get('categories', [])
    tot_exp = 0.0
    tot_cnt = 0

    for idx, c_item in enumerate(cats, 1):
        amt = float(c_item.get('amount', 0.0))
        cnt = int(c_item.get('count', 0))
        avg = round(amt / cnt, 2) if cnt > 0 else 0.0
        pct = c_item.get('percentage', 0.0)

        tot_exp += amt
        tot_cnt += cnt

        row_num = 4 + idx
        ws1.row_dimensions[row_num].height = 20
        row_vals = [idx, c_item.get('categoryName', '—'), amt, cnt, avg, f"{pct}%"]

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws1.cell(row=row_num, column=col_idx, value=val)
            cell.border = THIN_BORDER

            if col_idx in [3, 5]:
                cell.number_format = '₹#,##0'
                cell.alignment = Alignment(horizontal='right', vertical="center")
                cell.font = due_font if col_idx == 3 else bold_font
            elif col_idx in [1, 4, 6]:
                cell.font = normal_font
                cell.alignment = Alignment(horizontal='center', vertical="center")
            else:
                cell.font = bold_font
                cell.alignment = Alignment(horizontal='left', vertical="center")

    # Summary Row
    overall_avg = round(tot_exp / tot_cnt, 2) if tot_cnt > 0 else 0.0
    tot_row_num = 4 + len(cats) + 1
    ws1.row_dimensions[tot_row_num].height = 24

    sum_vals = ['TOTAL', 'Consolidated Outflows', tot_exp, tot_cnt, overall_avg, '100%']
    for col_idx, val in enumerate(sum_vals, 1):
        cell = ws1.cell(row=tot_row_num, column=col_idx, value=val)
        cell.border = TOTALS_BORDER
        cell.fill = TOTALS_FILL

        if col_idx in [3, 5]:
            cell.number_format = '₹#,##0'
            cell.alignment = Alignment(horizontal='right', vertical="center")
            cell.font = totals_red if col_idx == 3 else totals_bold
        elif col_idx in [1, 2]:
            cell.font = totals_bold
            cell.alignment = Alignment(horizontal='left', vertical="center")
        else:
            cell.font = totals_bold
            cell.alignment = Alignment(horizontal='center', vertical="center")

    for col in ws1.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row in (1, 2, 3):
                continue
            v_str = str(cell.value or '')
            if len(v_str) > max_len:
                max_len = len(v_str)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # Sheet 2: Budget Distribution Insights
    ws2 = wb.create_sheet(title="Budget Insights")
    ws2.views.sheetView[0].showGridLines = True
    ws2.merge_cells("A1:D1")
    ws2["A1"] = "OPERATIONAL OVERHEAD ALLOCATION & BUDGET CONTROL"
    ws2["A1"].font = Font(name="Calibri", size=12, bold=True, color="1A2B4C")
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 26

    ws2.merge_cells("A2:D2")
    ws2["A2"] = f"Scope: {branch_label} • Reconciled on {now_str}"
    ws2["A2"].font = Font(name="Calibri", size=9, italic=True, color="64748B")
    ws2["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[2].height = 18

    ws2.row_dimensions[3].height = 8

    cat_insights_headers = ['Budget Head / Ratio', 'Actual Spend', 'Cost Share', 'Control Status']
    ws2.row_dimensions[4].height = 22
    for c_idx, h in enumerate(cat_insights_headers, 1):
        c = ws2.cell(row=4, column=c_idx, value=h)
        c.font = header_font
        c.fill = HEADER_FILL
        c.border = THIN_BORDER
        c.alignment = Alignment(horizontal="center" if c_idx in [1, 4] else "left", vertical="center")

    cat_rows = []
    for c_item in cats[:6]:
        amt = float(c_item.get('amount', 0.0))
        pct = c_item.get('percentage', 0.0)
        status = "Major Head" if pct >= 25 else ("Standard" if pct >= 10 else "Minor Head")
        cat_rows.append([c_item.get('categoryName', '—'), format_currency_inr(amt), f"{pct}%", status])

    for r_idx, r_data in enumerate(cat_rows, 5):
        ws2.row_dimensions[r_idx].height = 20
        for c_idx, val in enumerate(r_data, 1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=val)
            cell.border = THIN_BORDER
            cell.font = bold_font if c_idx in [1, 2] else normal_font
            cell.alignment = Alignment(horizontal="center" if c_idx in [3, 4] else "left", vertical="center")

    for col in ws2.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(c.value or '')) for c in col if c.row > 3) or 15
        ws2.column_dimensions[col_letter].width = max_len + 4

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


# =============================================================================
# 6. Expense Categories PDF Export (Landscape A4 - Fixed from Portrait)
# =============================================================================

def export_expense_categories_pdf(categories_data, options=None):
    """
    Generates Landscape A4 Expense Categories Summary PDF report matching official snapshot layout:
    - Institutional header with logo, title, and scope box
    - Landscape table with Red disbursed amounts (#DC2626)
    - Reconciled totals row
    - Two-pass page numbering
    """
    options = options or {}
    branch_label = options.get('branch_name') or 'All Campuses'

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=20,
        rightMargin=20,
        topMargin=20,
        bottomMargin=30,
    )

    styles = getSampleStyleSheet()

    hdr_style = ParagraphStyle(
        'CatHdr',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=1,
    )
    hdr_right_style = ParagraphStyle(
        'CatHdrR',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=2,
    )
    cell_normal = ParagraphStyle(
        'CatCellNorm',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor(COLOR_SLATE_TEXT),
    )
    cell_bold = ParagraphStyle(
        'CatCellBld',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor(COLOR_SLATE_TEXT),
    )
    cell_center = ParagraphStyle(
        'CatCellCenter',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor(COLOR_SLATE_TEXT),
        alignment=1,
    )
    cell_right = ParagraphStyle(
        'CatCellRight',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor(COLOR_SLATE_TEXT),
        alignment=2,
    )
    cell_red = ParagraphStyle(
        'CatCellRed',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor(COLOR_RED_PENDING),
        alignment=2,
    )

    story = []

    # 1. Executive Snapshot Header
    cats = categories_data.get('categories', [])
    tot_exp = float(categories_data.get('totalExpenses', 0.0))
    tot_cnt = int(categories_data.get('totalVouchers', 0))

    header_meta = {
        'institute_name': 'SUTAWANE CLASSES VAIJAPUR',
        'branch_name': branch_label,
        'records_count': len(cats),
    }
    story.extend(build_executive_pdf_header(
        doc_title="OPERATIONAL EXPENSE CATEGORY DISTRIBUTION",
        meta=header_meta,
        report_badge_label="DISBURSED:",
        report_badge_value=format_currency_inr(tot_exp),
    ))

    # 2. Table Headers
    tbl_headers = [
        Paragraph("#", hdr_style),
        Paragraph("Expense Category", hdr_style),
        Paragraph("Total Disbursed (Rs.)", hdr_right_style),
        Paragraph("Vouchers Count", hdr_style),
        Paragraph("Average / Voucher (Rs.)", hdr_right_style),
        Paragraph("Budget Share (%)", hdr_style),
    ]
    tbl_rows = [tbl_headers]

    for idx, c_item in enumerate(cats, 1):
        amt = float(c_item.get('amount', 0.0))
        cnt = int(c_item.get('count', 0))
        avg = round(amt / cnt, 2) if cnt > 0 else 0.0
        pct = c_item.get('percentage', 0.0)

        tbl_rows.append([
            Paragraph(str(idx), cell_center),
            Paragraph(str(c_item.get('categoryName', '—')), cell_bold),
            Paragraph(format_currency_inr(amt), cell_red),
            Paragraph(str(cnt), cell_center),
            Paragraph(format_currency_inr(avg), cell_right),
            Paragraph(f"{pct}%", cell_center),
        ])

    overall_avg = round(tot_exp / tot_cnt, 2) if tot_cnt > 0 else 0.0
    tbl_rows.append([
        Paragraph("<b>TOTAL</b>", cell_bold),
        Paragraph("<b>Consolidated Outflows</b>", cell_bold),
        Paragraph(format_currency_inr(tot_exp), cell_red),
        Paragraph(f"<b>{tot_cnt}</b>", cell_center),
        Paragraph(format_currency_inr(overall_avg), cell_bold),
        Paragraph("<b>100%</b>", cell_center),
    ])

    # Col widths totaling 802 pt (Landscape A4)
    # [30, 222, 160, 110, 150, 130] = 802
    cat_tbl = Table(tbl_rows, colWidths=[30, 222, 160, 110, 150, 130])
    cat_tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(COLOR_NAVY_SECONDARY)),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor(COLOR_BORDER)),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor(COLOR_BG_TOTALS)),
        ('LINEABOVE', (0, -1), (-1, -1), 1.2, colors.HexColor(COLOR_NAVY_PRIMARY)),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#F8FAFC')]),
    ]))
    story.append(cat_tbl)

    doc.build(story, canvasmaker=ExecutiveLandscapeCanvas)
    buffer.seek(0)
    return buffer.getvalue()


# =============================================================================
# 7. Day Book / Financial Statement Excel Export (Multi-Sheet)
# =============================================================================

def export_statement_excel(statement_data, options=None):
    """
    Exports unified financial statement / Day Book to an Excel (.xlsx) workbook with:
    - Sheet 1: Executive Summary
    - Sheet 2: Fee Collections Ledger (Green #16A34A amounts)
    - Sheet 3: Center Disbursements Ledger (Red #DC2626 amounts)
    - Row heights: Title=28, Meta=18, Header=24, Rows=20, Totals=24
    - Auto column widths
    """
    options = options or {}
    period_label = options.get('period_label') or 'Financial Statement'
    branch_label = options.get('branch_name') or 'All Branches'
    session_label = options.get('session_name') or 'All Sessions'
    now_str = datetime.now().strftime("%d-%b-%Y %I:%M %p")

    wb = Workbook()

    bold_font = Font(name="Calibri", size=10, bold=True)
    normal_font = Font(name="Calibri", size=10)
    paid_font = Font(name="Calibri", size=10, bold=True, color="16A34A")
    due_font = Font(name="Calibri", size=10, bold=True, color="DC2626")
    totals_bold = Font(name="Calibri", size=11, bold=True, color="1A2B4C")
    totals_paid = Font(name="Calibri", size=11, bold=True, color="16A34A")
    totals_due = Font(name="Calibri", size=11, bold=True, color="DC2626")

    summary = statement_data.get('summary', {})
    tot_col = float(summary.get('totalCollections', 0.0))
    cnt_col = int(summary.get('collectionsCount', 0))
    tot_exp = float(summary.get('totalExpenses', 0.0))
    cnt_exp = int(summary.get('expensesCount', 0))
    net = tot_col - tot_exp

    # -------------------------------------------------------------
    # Sheet 1: Executive Summary
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Executive Summary"
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells("A1:D1")
    ws1["A1"] = "SUTAWANE CLASSES — UNIFIED DAY BOOK & FINANCIAL STATEMENT"
    ws1["A1"].font = Font(name="Calibri", size=13, bold=True, color="1A2B4C")
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 28

    ws1.merge_cells("A2:D2")
    ws1["A2"] = f"Period: {period_label} | Branch: {branch_label} | Session: {session_label}"
    ws1["A2"].font = Font(name="Calibri", size=9.5, italic=True, color="64748B")
    ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 18

    ws1.row_dimensions[3].height = 8

    sum_hdrs = ['Financial Metric Head', 'Realized Value (Rs.)', 'Voucher Volume', 'Audit Notes']
    ws1.row_dimensions[4].height = 24
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    for col_idx, h in enumerate(sum_hdrs, 1):
        c = ws1.cell(row=4, column=col_idx, value=h)
        c.font = header_font
        c.fill = HEADER_FILL
        c.border = THIN_BORDER
        c.alignment = Alignment(horizontal="center" if col_idx in [1, 3] else ("right" if col_idx == 2 else "left"), vertical="center")

    summary_rows = [
        ['Total Fee Collections', tot_col, f'{cnt_col} Receipts', 'Confirmed fees deposited into center accounts'],
        ['Total Center Expenditures', tot_exp, f'{cnt_exp} Vouchers', 'Operational disbursements & branch overheads'],
        ['Net Cash Flow / Operating Surplus', net, f'{cnt_col + cnt_exp} Transactions', 'Operating Surplus' if net >= 0 else 'Operating Deficit'],
    ]

    for idx, r_data in enumerate(summary_rows, 1):
        row_num = 4 + idx
        ws1.row_dimensions[row_num].height = 22
        for c_idx, val in enumerate(r_data, 1):
            cell = ws1.cell(row=row_num, column=c_idx, value=val)
            cell.border = THIN_BORDER
            if c_idx == 2:
                cell.number_format = '₹#,##0'
                cell.alignment = Alignment(horizontal='right', vertical="center")
                if idx == 1:
                    cell.font = paid_font
                elif idx == 2:
                    cell.font = due_font
                else:
                    cell.font = paid_font if net >= 0 else due_font
            elif c_idx in [1, 3]:
                cell.font = bold_font
                cell.alignment = Alignment(horizontal='center' if c_idx == 3 else 'left', vertical="center")
            else:
                cell.font = normal_font
                cell.alignment = Alignment(horizontal='left', vertical="center")

    for col in ws1.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(c.value or '')) for c in col if c.row > 3) or 15
        ws1.column_dimensions[col_letter].width = max(max_len + 5, 20)

    # -------------------------------------------------------------
    # Sheet 2: Fee Collections Ledger
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Fee Collections")
    ws2.views.sheetView[0].showGridLines = True

    col_headers = [
        '#', 'Receipt No.', 'Student Name', 'Roll No.', 'Class / Course',
        'Campus Branch', 'Payment Mode', 'Txn / Ref No.', 'Payment Date', 'Amount (Rs.)'
    ]
    max_col_letter2 = get_column_letter(len(col_headers))

    ws2.merge_cells(f"A1:{max_col_letter2}1")
    ws2["A1"] = f"FEE COLLECTIONS REGISTER — {period_label.upper()}"
    ws2["A1"].font = Font(name="Calibri", size=13, bold=True, color="1A2B4C")
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 28

    ws2.merge_cells(f"A2:{max_col_letter2}2")
    ws2["A2"] = f"Branch: {branch_label} | Generated on: {now_str}"
    ws2["A2"].font = Font(name="Calibri", size=9.5, italic=True, color="64748B")
    ws2["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[2].height = 18

    ws2.row_dimensions[3].height = 8

    ws2.row_dimensions[4].height = 24
    for col_idx, h in enumerate(col_headers, 1):
        cell = ws2.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        align_h = 'center' if col_idx in [1, 4, 7, 9] else ('right' if col_idx == 10 else 'left')
        cell.alignment = Alignment(horizontal=align_h, vertical="center")

    collections = statement_data.get('collections', [])
    tot_col_amt = 0.0

    for idx, p in enumerate(collections, 1):
        amt = float(p.get('totalAmount', 0.0))
        tot_col_amt += amt
        p_date = p.get('paymentDate', '')[:10] if p.get('paymentDate') else '—'

        row_num = 4 + idx
        ws2.row_dimensions[row_num].height = 20
        row_vals = [
            idx,
            p.get('receiptNumber', '—'),
            p.get('studentName', '—'),
            p.get('rollNumber', '—'),
            p.get('standard', '—'),
            p.get('branchName', '—'),
            p.get('paymentMode', '—'),
            p.get('referenceNumber', '—'),
            p_date,
            amt
        ]

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws2.cell(row=row_num, column=col_idx, value=val)
            cell.border = THIN_BORDER

            if col_idx == 10:
                cell.number_format = '₹#,##0'
                cell.alignment = Alignment(horizontal='right', vertical="center")
                cell.font = paid_font
            elif col_idx in [1, 4, 7, 9]:
                cell.font = normal_font
                cell.alignment = Alignment(horizontal='center', vertical="center")
            elif col_idx in [2, 3]:
                cell.font = bold_font
                cell.alignment = Alignment(horizontal='left', vertical="center")
            else:
                cell.font = normal_font
                cell.alignment = Alignment(horizontal='left', vertical="center")

    # Collections Totals Row
    tot_row_num2 = 4 + len(collections) + 1
    ws2.row_dimensions[tot_row_num2].height = 24
    sum_vals_col = ['TOTAL', f'{len(collections)} Receipts', '', '', '', '', '', '', '', tot_col_amt]

    for col_idx, val in enumerate(sum_vals_col, 1):
        cell = ws2.cell(row=tot_row_num2, column=col_idx, value=val)
        cell.border = TOTALS_BORDER
        cell.fill = TOTALS_FILL

        if col_idx == 10:
            cell.number_format = '₹#,##0'
            cell.alignment = Alignment(horizontal='right', vertical="center")
            cell.font = totals_paid
        elif col_idx in [1, 2]:
            cell.font = totals_bold
            cell.alignment = Alignment(horizontal='left', vertical="center")
        else:
            cell.alignment = Alignment(horizontal='center', vertical="center")

    for col in ws2.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row in (1, 2, 3):
                continue
            v_str = str(cell.value or '')
            if len(v_str) > max_len:
                max_len = len(v_str)
        ws2.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # -------------------------------------------------------------
    # Sheet 3: Center Disbursements Ledger
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Center Disbursements")
    ws3.views.sheetView[0].showGridLines = True

    exp_headers = [
        '#', 'Voucher / Ref No.', 'Expense Category', 'Description / Purpose',
        'Paid To / Vendor', 'Campus Branch', 'Payment Mode', 'Disbursement Date', 'Amount (Rs.)'
    ]
    max_col_letter3 = get_column_letter(len(exp_headers))

    ws3.merge_cells(f"A1:{max_col_letter3}1")
    ws3["A1"] = f"CENTER EXPENDITURES REGISTER — {period_label.upper()}"
    ws3["A1"].font = Font(name="Calibri", size=13, bold=True, color="1A2B4C")
    ws3["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[1].height = 28

    ws3.merge_cells(f"A2:{max_col_letter3}2")
    ws3["A2"] = f"Branch: {branch_label} | Generated on: {now_str}"
    ws3["A2"].font = Font(name="Calibri", size=9.5, italic=True, color="64748B")
    ws3["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[2].height = 18

    ws3.row_dimensions[3].height = 8

    ws3.row_dimensions[4].height = 24
    for col_idx, h in enumerate(exp_headers, 1):
        cell = ws3.cell(row=4, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = HEADER_FILL
        cell.border = THIN_BORDER
        align_h = 'center' if col_idx in [1, 7, 8] else ('right' if col_idx == 9 else 'left')
        cell.alignment = Alignment(horizontal=align_h, vertical="center")

    expenses = statement_data.get('expenses', [])
    tot_exp_amt = 0.0

    for idx, e in enumerate(expenses, 1):
        amt = float(e.get('amount', 0.0))
        tot_exp_amt += amt
        e_date = e.get('expenseDate', '')[:10] if e.get('expenseDate') else '—'

        row_num = 4 + idx
        ws3.row_dimensions[row_num].height = 20
        row_vals = [
            idx,
            e.get('referenceNumber', '—'),
            e.get('categoryName', '—'),
            e.get('description', '—'),
            e.get('vendor') or e.get('paidTo') or '—',
            e.get('branchName', '—'),
            e.get('paymentMode', '—'),
            e_date,
            amt
        ]

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws3.cell(row=row_num, column=col_idx, value=val)
            cell.border = THIN_BORDER

            if col_idx == 9:
                cell.number_format = '₹#,##0'
                cell.alignment = Alignment(horizontal='right', vertical="center")
                cell.font = due_font
            elif col_idx in [1, 7, 8]:
                cell.font = normal_font
                cell.alignment = Alignment(horizontal='center', vertical="center")
            elif col_idx in [2, 3]:
                cell.font = bold_font
                cell.alignment = Alignment(horizontal='left', vertical="center")
            else:
                cell.font = normal_font
                cell.alignment = Alignment(horizontal='left', vertical="center")

    # Expenses Totals Row
    tot_row_num3 = 4 + len(expenses) + 1
    ws3.row_dimensions[tot_row_num3].height = 24
    sum_vals_exp = ['TOTAL', f'{len(expenses)} Vouchers', '', '', '', '', '', '', tot_exp_amt]

    for col_idx, val in enumerate(sum_vals_exp, 1):
        cell = ws3.cell(row=tot_row_num3, column=col_idx, value=val)
        cell.border = TOTALS_BORDER
        cell.fill = TOTALS_FILL

        if col_idx == 9:
            cell.number_format = '₹#,##0'
            cell.alignment = Alignment(horizontal='right', vertical="center")
            cell.font = totals_due
        elif col_idx in [1, 2]:
            cell.font = totals_bold
            cell.alignment = Alignment(horizontal='left', vertical="center")
        else:
            cell.alignment = Alignment(horizontal='center', vertical="center")

    for col in ws3.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            if cell.row in (1, 2, 3):
                continue
            v_str = str(cell.value or '')
            if len(v_str) > max_len:
                max_len = len(v_str)
        ws3.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


# =============================================================================
# 8. Day Book / Financial Statement PDF Export (Landscape A4)
# =============================================================================

def export_statement_pdf(statement_data, options=None):
    """
    Generates high-grade Landscape A4 Day Book / Statement PDF with:
    - Executive Snapshot header (matching user's snapshot)
    - KPI Metrics Summary strip
    - Dual tabular sections: Collections (Green amounts) & Expenses (Red amounts)
    - Reconciled totals and institutional signatures block
    """
    options = options or {}
    period_label = options.get('period_label') or 'Day Book Statement'
    branch_label = options.get('branch_name') or 'All Campuses (Consolidated)'
    session_label = options.get('session_name') or 'Academic Session'

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=20,
        rightMargin=20,
        topMargin=20,
        bottomMargin=30,
    )

    styles = getSampleStyleSheet()

    section_title = ParagraphStyle(
        'StmtSecTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9.5,
        textColor=colors.HexColor(COLOR_NAVY_PRIMARY),
        leading=13,
    )
    cell_normal = ParagraphStyle(
        'StmtCellNorm',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7,
        leading=9,
        textColor=colors.HexColor(COLOR_SLATE_TEXT),
    )
    cell_bold = ParagraphStyle(
        'StmtCellBld',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7,
        leading=9,
        textColor=colors.HexColor(COLOR_SLATE_TEXT),
    )
    cell_right = ParagraphStyle(
        'StmtCellRight',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7,
        leading=9,
        textColor=colors.HexColor(COLOR_SLATE_TEXT),
        alignment=2,
    )
    cell_center = ParagraphStyle(
        'StmtCellCenter',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7,
        leading=9,
        textColor=colors.HexColor(COLOR_SLATE_TEXT),
        alignment=1,
    )
    cell_green = ParagraphStyle(
        'StmtCellGreen',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7,
        leading=9,
        textColor=colors.HexColor(COLOR_GREEN_PAID),
        alignment=2,
    )
    cell_red = ParagraphStyle(
        'StmtCellRed',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7,
        leading=9,
        textColor=colors.HexColor(COLOR_RED_PENDING),
        alignment=2,
    )
    hdr_style = ParagraphStyle(
        'StmtHdr',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.white,
        alignment=1,
    )
    hdr_right_style = ParagraphStyle(
        'StmtHdrR',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.white,
        alignment=2,
    )

    story = []

    # 1. Executive Snapshot Header
    collections = statement_data.get('collections', [])
    expenses = statement_data.get('expenses', [])
    summary = statement_data.get('summary', {})
    tot_col = float(summary.get('totalCollections', 0.0))
    cnt_col = int(summary.get('collectionsCount', 0))
    tot_exp = float(summary.get('totalExpenses', 0.0))
    cnt_exp = int(summary.get('expensesCount', 0))
    net = tot_col - tot_exp

    header_meta = {
        'institute_name': 'SUTAWANE CLASSES VAIJAPUR',
        'branch_name': branch_label,
        'session_name': session_label,
        'records_count': len(collections) + len(expenses),
    }
    story.extend(build_executive_pdf_header(
        doc_title="OFFICIAL DAY BOOK & CASH FLOW STATEMENT",
        meta=header_meta,
        report_badge_label="NET CASH:",
        report_badge_value=format_currency_inr(net),
    ))

    # 2. Executive Metrics Summary Bar
    summary_bar_data = [
        [
            Paragraph(f"<font size=7 color='#047857'><b>TOTAL COLLECTIONS</b></font><br/><font size=10 color='#065F46'><b>{format_currency_inr(tot_col)}</b></font><br/><font size=6.5 color='#047857'>{cnt_col} Receipts Realized</font>", styles['Normal']),
            Paragraph(f"<font size=7 color='#BE123C'><b>CENTER EXPENSES</b></font><br/><font size=10 color='#9F1239'><b>{format_currency_inr(tot_exp)}</b></font><br/><font size=6.5 color='#BE123C'>{cnt_exp} Vouchers Disbursed</font>", styles['Normal']),
            Paragraph(f"<font size=7 color='#{'047857' if net >= 0 else 'BE123C'}'><b>NET CASH MOVEMENT</b></font><br/><font size=10 color='#{'065F46' if net >= 0 else '9F1239'}'><b>{format_currency_inr(net)}</b></font><br/><font size=6.5 color='#64748B'>{'Surplus Balance' if net >= 0 else 'Operating Deficit'}</font>", styles['Normal']),
            Paragraph(f"<font size=7 color='#1E40AF'><b>STATEMENT STATUS</b></font><br/><font size=10 color='#1E3A8A'><b>RECONCILED</b></font><br/><font size=6.5 color='#1E40AF'>Audit Trail Verified</font>", styles['Normal']),
        ]
    ]
    summary_tbl = Table(summary_bar_data, colWidths=[200, 200, 200, 202])
    summary_tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor(COLOR_BG_CARD)),
        ('BOX', (0, 0), (-1, -1), 0.75, colors.HexColor(COLOR_BORDER)),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(summary_tbl)
    story.append(Spacer(1, 10))

    # 3. Fee Collections Table
    story.append(Paragraph(f"<b>Part A: Fee Inflow Ledger ({len(collections)} Receipts Realized)</b>", section_title))
    story.append(Spacer(1, 4))

    col_tbl_headers = [
        Paragraph("#", hdr_style),
        Paragraph("Receipt No", hdr_style),
        Paragraph("Student Name", hdr_style),
        Paragraph("Roll No", hdr_style),
        Paragraph("Class / Course", hdr_style),
        Paragraph("Branch", hdr_style),
        Paragraph("Mode", hdr_style),
        Paragraph("Reference / Txn", hdr_style),
        Paragraph("Date", hdr_style),
        Paragraph("Amount (Rs.)", hdr_right_style),
    ]
    col_tbl_rows = [col_tbl_headers]

    for idx, p in enumerate(collections, 1):
        p_date = p.get('paymentDate', '')[:10] if p.get('paymentDate') else '—'
        amt = float(p.get('totalAmount', 0.0))
        col_tbl_rows.append([
            Paragraph(str(idx), cell_center),
            Paragraph(str(p.get('receiptNumber', '—')), cell_bold),
            Paragraph(str(p.get('studentName', '—')), cell_bold),
            Paragraph(str(p.get('rollNumber', '—')), cell_center),
            Paragraph(str(p.get('standard', '—')), cell_normal),
            Paragraph(str(p.get('branchName', '—')), cell_normal),
            Paragraph(str(p.get('paymentMode', '—')), cell_center),
            Paragraph(str(p.get('referenceNumber') or '—'), cell_normal),
            Paragraph(p_date, cell_center),
            Paragraph(format_currency_inr(amt), cell_green),
        ])

    col_tbl_rows.append([
        Paragraph("<b>TOTAL</b>", cell_bold),
        Paragraph(f"<b>{len(collections)} Receipts</b>", cell_bold),
        Paragraph("", cell_normal),
        Paragraph("", cell_normal),
        Paragraph("", cell_normal),
        Paragraph("", cell_normal),
        Paragraph("", cell_normal),
        Paragraph("", cell_normal),
        Paragraph("", cell_normal),
        Paragraph(format_currency_inr(tot_col), cell_green),
    ])

    # Col widths totaling 802 pt (Landscape A4)
    # [24, 75, 140, 55, 80, 80, 60, 110, 68, 110] = 802
    col_table = Table(col_tbl_rows, colWidths=[24, 75, 140, 55, 80, 80, 60, 110, 68, 110])
    col_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(COLOR_NAVY_SECONDARY)),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor(COLOR_BORDER)),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor(COLOR_BG_TOTALS)),
        ('LINEABOVE', (0, -1), (-1, -1), 1.2, colors.HexColor(COLOR_NAVY_PRIMARY)),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#F8FAFC')]),
    ]))
    story.append(col_table)
    story.append(Spacer(1, 12))

    # 4. Center Expenses Table
    story.append(Paragraph(f"<b>Part B: Center Disbursements Ledger ({len(expenses)} Vouchers Incurred)</b>", section_title))
    story.append(Spacer(1, 4))

    exp_tbl_headers = [
        Paragraph("#", hdr_style),
        Paragraph("Voucher / Ref", hdr_style),
        Paragraph("Category", hdr_style),
        Paragraph("Description / Purpose", hdr_style),
        Paragraph("Paid To / Vendor", hdr_style),
        Paragraph("Branch", hdr_style),
        Paragraph("Mode", hdr_style),
        Paragraph("Date", hdr_style),
        Paragraph("Amount (Rs.)", hdr_right_style),
    ]
    exp_tbl_rows = [exp_tbl_headers]

    for idx, e in enumerate(expenses, 1):
        e_date = e.get('expenseDate', '')[:10] if e.get('expenseDate') else '—'
        amt = float(e.get('amount', 0.0))
        exp_tbl_rows.append([
            Paragraph(str(idx), cell_center),
            Paragraph(str(e.get('referenceNumber') or f"EXP-{idx:04d}"), cell_bold),
            Paragraph(str(e.get('categoryName', '—')), cell_bold),
            Paragraph(str(e.get('description', '—')), cell_normal),
            Paragraph(str(e.get('vendor') or e.get('paidTo') or '—'), cell_normal),
            Paragraph(str(e.get('branchName', '—')), cell_normal),
            Paragraph(str(e.get('paymentMode', '—')), cell_center),
            Paragraph(e_date, cell_center),
            Paragraph(format_currency_inr(amt), cell_red),
        ])

    exp_tbl_rows.append([
        Paragraph("<b>TOTAL</b>", cell_bold),
        Paragraph(f"<b>{len(expenses)} Vouchers</b>", cell_bold),
        Paragraph("", cell_normal),
        Paragraph("", cell_normal),
        Paragraph("", cell_normal),
        Paragraph("", cell_normal),
        Paragraph("", cell_normal),
        Paragraph("", cell_normal),
        Paragraph(format_currency_inr(tot_exp), cell_red),
    ])

    # Col widths totaling 802 pt (Landscape A4)
    # [24, 85, 105, 175, 110, 75, 60, 68, 100] = 802
    exp_table = Table(exp_tbl_rows, colWidths=[24, 85, 105, 175, 110, 75, 60, 68, 100])
    exp_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(COLOR_NAVY_SECONDARY)),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor(COLOR_BORDER)),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor(COLOR_BG_TOTALS)),
        ('LINEABOVE', (0, -1), (-1, -1), 1.2, colors.HexColor(COLOR_NAVY_PRIMARY)),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#F8FAFC')]),
    ]))
    story.append(exp_table)
    story.append(Spacer(1, 14))

    # 5. Signatures Block
    sig_data = [
        [
            Paragraph("Prepared By:<br/><br/>_______________________<br/>Accounts Executive", cell_normal),
            Paragraph("Verified By:<br/><br/>_______________________<br/>Center Manager", cell_normal),
            Paragraph("Approved &amp; Audited By:<br/><br/>_______________________<br/>Managing Director / Owner", cell_right),
        ]
    ]
    sig_tbl = Table(sig_data, colWidths=[264, 264, 274])
    sig_tbl.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(sig_tbl)

    doc.build(story, canvasmaker=ExecutiveLandscapeCanvas)
    buffer.seek(0)
    return buffer.getvalue()
