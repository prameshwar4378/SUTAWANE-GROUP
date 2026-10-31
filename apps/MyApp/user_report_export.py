"""
User Activity & Staff Operations Audit Report Export Service.
Generates comprehensive multi-sheet Openpyxl Excel workbooks and
publication-grade Landscape A4 ReportLab PDFs for auditing individual staff operations.
"""

import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.platypus import Paragraph, Spacer, Table, TableStyle, SimpleDocTemplate
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from .export_theme import (
    COLOR_NAVY_PRIMARY, COLOR_NAVY_SECONDARY,
    COLOR_GREEN_PAID, COLOR_RED_PENDING, COLOR_BORDER, COLOR_BG_CARD,
    COLOR_BG_TOTALS, COLOR_SLATE_TEXT,
    FONT_FAMILY, FONT_FAMILY_BOLD,
    ExecutiveLandscapeCanvas, build_executive_pdf_header,
    format_currency_inr
)


# =============================================================================
# EXCEL EXPORT SERVICE
# =============================================================================

def export_user_report_excel(report_data, meta=None):
    """
    Generates an executive multi-sheet Excel workbook for User Activity Audit:
    - Sheet 1: Executive Overview & KPI Summary
    - Sheet 2: Fee Collections Realized
    - Sheet 3: Center Disbursements & Expenses
    - Sheet 4: Student Admissions Enrolled
    - Sheet 5: Fee Discounts Authorized
    """
    meta = meta or {}
    user_info = report_data.get('user') or {}
    user_name = user_info.get('name') or 'All Staff Members Combined'
    user_role = user_info.get('role_display') or 'All Roles'
    period_label = report_data.get('period', {}).get('label') or 'All Time'
    branch_label = meta.get('branch_label') or report_data.get('branch_label') or 'All Branches'
    session_label = meta.get('session_label') or report_data.get('session_label') or 'All Sessions'
    summary = report_data.get('summary', {})
    now_str = datetime.now().strftime("%d-%b-%Y %I:%M %p")

    wb = Workbook()
    wb.remove(wb.active)  # Remove default sheet

    # Fonts & Styles
    title_font = Font(name="Calibri", size=14, bold=True, color="1A2B4C")
    meta_font = Font(name="Calibri", size=9.5, italic=True, color="64748B")
    sec_font = Font(name="Calibri", size=11, bold=True, color="1A2B4C")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=10, bold=True)
    normal_font = Font(name="Calibri", size=9.5)
    paid_font = Font(name="Calibri", size=9.5, bold=True, color="047857")
    outflow_font = Font(name="Calibri", size=9.5, bold=True, color="BE123C")
    totals_bold = Font(name="Calibri", size=10.5, bold=True, color="1A2B4C")

    HEADER_FILL = PatternFill(start_color="28375C", end_color="28375C", fill_type="solid")
    TOTALS_FILL = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    THIN_BORDER = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    TOTALS_BORDER = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='medium', color='1A2B4C'),
        bottom=Side(style='double', color='1A2B4C')
    )

    # -------------------------------------------------------------------------
    # SHEET 1: EXECUTIVE OVERVIEW & KPI SUMMARY
    # -------------------------------------------------------------------------
    ws1 = wb.create_sheet(title="Executive Overview")
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells("A1:G1")
    ws1["A1"] = f"USER & STAFF OPERATIONS AUDIT REPORT — {user_name.upper()}"
    ws1["A1"].font = title_font
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 28

    ws1.merge_cells("A2:G2")
    ws1["A2"] = f"Staff Member: {user_name} ({user_role}) | Scope: {branch_label} | Session: {session_label} | Period: {period_label} | Generated: {now_str}"
    ws1["A2"].font = meta_font
    ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 18

    # KPI Summary Cards Block
    ws1.row_dimensions[4].height = 20
    ws1["A4"] = "OPERATIONAL PERFORMANCE METRICS"
    ws1["A4"].font = sec_font

    kpi_headers = ["Metric Description", "Volume / Count", "Realized Amount (₹)", "Status / Notes"]
    ws1.row_dimensions[5].height = 24
    for idx, kh in enumerate(kpi_headers, start=1):
        c = ws1.cell(row=5, column=idx, value=kh)
        c.font = header_font
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center" if idx in [2, 4] else ("right" if idx == 3 else "left"), vertical="center")
        c.border = THIN_BORDER

    tot_col = float(summary.get('total_collected', 0.0))
    cnt_col = int(summary.get('payments_count', 0))
    tot_exp = float(summary.get('total_expenses', 0.0))
    cnt_exp = int(summary.get('expenses_count', 0))
    net_val = tot_col - tot_exp
    cnt_adm = int(summary.get('admissions_count', 0))
    tot_adm_fee = float(summary.get('total_agreed_fees', 0.0))
    cnt_disc = int(summary.get('discounts_count', 0))
    tot_disc = float(summary.get('total_discounts', 0.0))

    kpi_rows = [
        ("Fee Payments Realized", f"{cnt_col} receipts", tot_col, "Active Inflows"),
        ("Center Operational Disbursements", f"{cnt_exp} vouchers", tot_exp, "Approved Expenses"),
        ("Net Operating Cash Movement", "Inflow - Outflow", net_val, "Surplus Balance" if net_val >= 0 else "Deficit"),
        ("New Student Admissions Registered", f"{cnt_adm} students", tot_adm_fee, "Agreed Fee Volume"),
        ("Fee Concessions & Discounts Granted", f"{cnt_disc} concessions", tot_disc, "Authorized Concessions"),
    ]

    for r_idx, (m_desc, m_cnt, m_val, m_note) in enumerate(kpi_rows, start=6):
        ws1.row_dimensions[r_idx].height = 20
        c1 = ws1.cell(row=r_idx, column=1, value=m_desc)
        c1.font = bold_font
        c1.border = THIN_BORDER

        c2 = ws1.cell(row=r_idx, column=2, value=m_cnt)
        c2.font = normal_font
        c2.alignment = Alignment(horizontal="center", vertical="center")
        c2.border = THIN_BORDER

        c3 = ws1.cell(row=r_idx, column=3, value=m_val)
        c3.number_format = '₹#,##0.00'
        c3.alignment = Alignment(horizontal="right", vertical="center")
        c3.font = paid_font if r_idx in [6, 8] and m_val >= 0 else (outflow_font if r_idx in [7, 10] or m_val < 0 else bold_font)
        c3.border = THIN_BORDER

        c4 = ws1.cell(row=r_idx, column=4, value=m_note)
        c4.font = normal_font
        c4.alignment = Alignment(horizontal="center", vertical="center")
        c4.border = THIN_BORDER

    # Mode-wise Breakdown Table
    ws1.row_dimensions[12].height = 20
    ws1["A12"] = "COLLECTIONS BREAKDOWN BY PAYMENT CHANNEL"
    ws1["A12"].font = sec_font

    mode_headers = ["#", "Payment Mode / Channel", "Receipts Count", "Amount Collected (₹)", "Share (%)"]
    ws1.row_dimensions[13].height = 22
    for idx, mh in enumerate(mode_headers, start=1):
        c = ws1.cell(row=13, column=idx, value=mh)
        c.font = header_font
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center" if idx in [1, 3, 5] else ("right" if idx == 4 else "left"), vertical="center")
        c.border = THIN_BORDER

    collections_by_mode = summary.get('collections_by_mode', [])
    m_row = 14
    for idx, mb in enumerate(collections_by_mode, start=1):
        ws1.row_dimensions[m_row].height = 19
        pct = (float(mb['total']) / tot_col * 100) if tot_col > 0 else 0.0
        row_vals = [idx, mb['mode'], mb['count'], float(mb['total']), f"{pct:.1f}%"]
        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws1.cell(row=m_row, column=c_idx, value=val)
            cell.font = normal_font
            cell.border = THIN_BORDER
            if c_idx in [1, 3, 5]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif c_idx == 4:
                cell.number_format = '₹#,##0.00'
                cell.font = paid_font
                cell.alignment = Alignment(horizontal="right", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        m_row += 1

    for col in ws1.columns:
        max_len = max(len(str(c.value or '')) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # -------------------------------------------------------------------------
    # SHEET 2: FEE COLLECTIONS REGISTER
    # -------------------------------------------------------------------------
    ws2 = wb.create_sheet(title="Fee Collections")
    ws2.views.sheetView[0].showGridLines = True

    col_headers = [
        '#', 'Receipt No.', 'Payment Date', 'Student Code', 'Student Name', 'Roll No.',
        'Class / Course', 'Campus Branch', 'Academic Session', 'Payment Mode',
        'Amount Paid (Rs.)', 'Payment Status', 'Received By', 'Initiated By'
    ]
    max_col2 = get_column_letter(len(col_headers))

    ws2.merge_cells(f"A1:{max_col2}1")
    ws2["A1"] = f"FEE COLLECTIONS REGISTER — {user_name.upper()}"
    ws2["A1"].font = title_font
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 28

    ws2.merge_cells(f"A2:{max_col2}2")
    ws2["A2"] = f"Period: {period_label} | Total Receipts: {len(report_data.get('collections', []))} | Scope: {branch_label}"
    ws2["A2"].font = meta_font
    ws2["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[2].height = 18

    ws2.row_dimensions[4].height = 24
    for idx, h in enumerate(col_headers, start=1):
        c = ws2.cell(row=4, column=idx, value=h)
        c.font = header_font
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center" if idx in [1, 3, 4, 6, 10, 12] else ("right" if idx == 11 else "left"), vertical="center")
        c.border = THIN_BORDER

    cur_row = 5
    tot_col_sheet = 0.0
    for idx, p in enumerate(report_data.get('collections', []), start=1):
        amt = float(p.get('amount', 0.0))
        tot_col_sheet += amt
        ws2.row_dimensions[cur_row].height = 20

        row_vals = [
            idx,
            p.get('receipt_number', '—'),
            p.get('payment_date', '—'),
            p.get('student_code', '—'),
            p.get('student_name', '—'),
            p.get('roll_number', '—'),
            p.get('standard', '—'),
            p.get('branch_name', '—'),
            p.get('session_name', '—'),
            p.get('payment_mode', '—'),
            amt,
            p.get('status', 'ACTIVE'),
            p.get('received_by', 'Staff'),
            p.get('created_by', 'Staff'),
        ]

        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws2.cell(row=cur_row, column=c_idx, value=val)
            cell.font = normal_font
            cell.border = THIN_BORDER
            if c_idx == 11:
                cell.number_format = '₹#,##0.00'
                cell.font = paid_font
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif c_idx in [1, 3, 4, 6, 10, 12]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                if c_idx in [2, 5]:
                    cell.font = bold_font
        cur_row += 1

    # Totals Row
    ws2.row_dimensions[cur_row].height = 24
    ws2.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=10)
    lbl = ws2.cell(row=cur_row, column=1, value=f"TOTAL REALIZED INFLOWS ({len(report_data.get('collections', []))} Receipts)")
    lbl.font = totals_bold
    lbl.alignment = Alignment(horizontal="right", vertical="center")

    tot_cell = ws2.cell(row=cur_row, column=11, value=tot_col_sheet)
    tot_cell.font = totals_bold
    tot_cell.number_format = '₹#,##0.00'
    tot_cell.alignment = Alignment(horizontal="right", vertical="center")

    for col in range(1, len(col_headers) + 1):
        c = ws2.cell(row=cur_row, column=col)
        c.border = TOTALS_BORDER
        c.fill = TOTALS_FILL

    for col in ws2.columns:
        max_len = max(len(str(c.value or '')) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # -------------------------------------------------------------------------
    # SHEET 3: CENTER DISBURSEMENTS & EXPENSES
    # -------------------------------------------------------------------------
    ws3 = wb.create_sheet(title="Center Expenses")
    ws3.views.sheetView[0].showGridLines = True

    exp_headers = [
        '#', 'Voucher / Ref No.', 'Disbursement Date', 'Expense Category', 'Voucher Purpose / Description',
        'Paid To / Vendor', 'Branch Location', 'Payment Mode', 'Outflow Amount (Rs.)', 'Recorded By'
    ]
    max_col3 = get_column_letter(len(exp_headers))

    ws3.merge_cells(f"A1:{max_col3}1")
    ws3["A1"] = f"CENTER EXPENDITURES REGISTER — {user_name.upper()}"
    ws3["A1"].font = title_font
    ws3["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[1].height = 28

    ws3.merge_cells(f"A2:{max_col3}2")
    ws3["A2"] = f"Period: {period_label} | Total Vouchers: {len(report_data.get('expenses', []))} | Scope: {branch_label}"
    ws3["A2"].font = meta_font
    ws3["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[2].height = 18

    ws3.row_dimensions[4].height = 24
    for idx, h in enumerate(exp_headers, start=1):
        c = ws3.cell(row=4, column=idx, value=h)
        c.font = header_font
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center" if idx in [1, 3, 8] else ("right" if idx == 9 else "left"), vertical="center")
        c.border = THIN_BORDER

    cur_row = 5
    tot_exp_sheet = 0.0
    for idx, e in enumerate(report_data.get('expenses', []), start=1):
        amt = float(e.get('amount', 0.0))
        tot_exp_sheet += amt
        ws3.row_dimensions[cur_row].height = 20

        row_vals = [
            idx,
            e.get('reference_number', '—'),
            e.get('expense_date', '—'),
            e.get('category_name', 'General'),
            e.get('description', '—'),
            e.get('vendor', '—'),
            e.get('branch_name', '—'),
            e.get('payment_mode', '—'),
            amt,
            e.get('recorded_by', 'Staff'),
        ]

        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws3.cell(row=cur_row, column=c_idx, value=val)
            cell.font = normal_font
            cell.border = THIN_BORDER
            if c_idx == 9:
                cell.number_format = '₹#,##0.00'
                cell.font = outflow_font
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif c_idx in [1, 3, 8]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                if c_idx in [2, 4]:
                    cell.font = bold_font
        cur_row += 1

    # Totals Row
    ws3.row_dimensions[cur_row].height = 24
    ws3.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=8)
    lbl = ws3.cell(row=cur_row, column=1, value=f"TOTAL DISBURSED OUTFLOWS ({len(report_data.get('expenses', []))} Vouchers)")
    lbl.font = totals_bold
    lbl.alignment = Alignment(horizontal="right", vertical="center")

    tot_cell = ws3.cell(row=cur_row, column=9, value=tot_exp_sheet)
    tot_cell.font = totals_bold
    tot_cell.number_format = '₹#,##0.00'
    tot_cell.alignment = Alignment(horizontal="right", vertical="center")

    for col in range(1, len(exp_headers) + 1):
        c = ws3.cell(row=cur_row, column=col)
        c.border = TOTALS_BORDER
        c.fill = TOTALS_FILL

    for col in ws3.columns:
        max_len = max(len(str(c.value or '')) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws3.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # -------------------------------------------------------------------------
    # SHEET 4: STUDENT ADMISSIONS ENROLLED
    # -------------------------------------------------------------------------
    ws4 = wb.create_sheet(title="Admissions Enrolled")
    ws4.views.sheetView[0].showGridLines = True

    adm_headers = [
        '#', 'Student Code', 'Roll No.', 'Student Full Name', 'Parent Contact',
        'Standard / Course', 'Campus Branch', 'Academic Session', 'Admission Date',
        'Agreed Total Fee (Rs.)', 'Discount Granted (Rs.)', 'Net Payable (Rs.)', 'Configured By'
    ]
    max_col4 = get_column_letter(len(adm_headers))

    ws4.merge_cells(f"A1:{max_col4}1")
    ws4["A1"] = f"STUDENT ADMISSIONS & ENROLLMENTS REGISTER — {user_name.upper()}"
    ws4["A1"].font = title_font
    ws4["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws4.row_dimensions[1].height = 28

    ws4.merge_cells(f"A2:{max_col4}2")
    ws4["A2"] = f"Period: {period_label} | Total Students Enrolled: {len(report_data.get('admissions', []))} | Scope: {branch_label}"
    ws4["A2"].font = meta_font
    ws4["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws4.row_dimensions[2].height = 18

    ws4.row_dimensions[4].height = 24
    for idx, h in enumerate(adm_headers, start=1):
        c = ws4.cell(row=4, column=idx, value=h)
        c.font = header_font
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center" if idx in [1, 2, 3, 5, 9] else ("right" if idx in [10, 11, 12] else "left"), vertical="center")
        c.border = THIN_BORDER

    cur_row = 5
    tot_agr = 0.0
    tot_dsc = 0.0
    tot_net = 0.0
    for idx, a in enumerate(report_data.get('admissions', []), start=1):
        t_fee = float(a.get('total_fees', 0.0))
        d_fee = float(a.get('total_discount', 0.0))
        n_fee = float(a.get('net_payable', 0.0))
        tot_agr += t_fee
        tot_dsc += d_fee
        tot_net += n_fee
        ws4.row_dimensions[cur_row].height = 20

        row_vals = [
            idx,
            a.get('student_code', '—'),
            a.get('roll_number', '—'),
            a.get('student_name', '—'),
            a.get('parent_phone', '—'),
            a.get('standard', '—'),
            a.get('branch_name', '—'),
            a.get('session_name', '—'),
            a.get('admission_date', '—'),
            t_fee,
            d_fee,
            n_fee,
            a.get('enrolled_by', 'Staff'),
        ]

        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws4.cell(row=cur_row, column=c_idx, value=val)
            cell.font = normal_font
            cell.border = THIN_BORDER
            if c_idx in [10, 11, 12]:
                cell.number_format = '₹#,##0.00'
                cell.font = bold_font
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif c_idx in [1, 2, 3, 5, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                if c_idx in [2, 4]:
                    cell.font = bold_font
        cur_row += 1

    # Totals Row
    ws4.row_dimensions[cur_row].height = 24
    ws4.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=9)
    lbl = ws4.cell(row=cur_row, column=1, value=f"TOTAL ENROLLMENT VOLUME ({len(report_data.get('admissions', []))} Students)")
    lbl.font = totals_bold
    lbl.alignment = Alignment(horizontal="right", vertical="center")

    c10 = ws4.cell(row=cur_row, column=10, value=tot_agr)
    c10.font = totals_bold
    c10.number_format = '₹#,##0.00'
    c10.alignment = Alignment(horizontal="right", vertical="center")

    c11 = ws4.cell(row=cur_row, column=11, value=tot_dsc)
    c11.font = totals_bold
    c11.number_format = '₹#,##0.00'
    c11.alignment = Alignment(horizontal="right", vertical="center")

    c12 = ws4.cell(row=cur_row, column=12, value=tot_net)
    c12.font = totals_bold
    c12.number_format = '₹#,##0.00'
    c12.alignment = Alignment(horizontal="right", vertical="center")

    for col in range(1, len(adm_headers) + 1):
        c = ws4.cell(row=cur_row, column=col)
        c.border = TOTALS_BORDER
        c.fill = TOTALS_FILL

    for col in ws4.columns:
        max_len = max(len(str(c.value or '')) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws4.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # -------------------------------------------------------------------------
    # SHEET 5: FEE DISCOUNTS AUTHORIZED
    # -------------------------------------------------------------------------
    ws5 = wb.create_sheet(title="Discounts Authorized")
    ws5.views.sheetView[0].showGridLines = True

    dsc_headers = [
        '#', 'Student Full Name', 'Roll No.', 'Standard / Course', 'Campus Branch',
        'Discount Reason / Concession Label', 'Concession Amount (Rs.)', 'Granted Date', 'Authorized By'
    ]
    max_col5 = get_column_letter(len(dsc_headers))

    ws5.merge_cells(f"A1:{max_col5}1")
    ws5["A1"] = f"FEE CONCESSIONS & DISCOUNTS AUTHORIZED — {user_name.upper()}"
    ws5["A1"].font = title_font
    ws5["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws5.row_dimensions[1].height = 28

    ws5.merge_cells(f"A2:{max_col5}2")
    ws5["A2"] = f"Period: {period_label} | Total Concessions: {len(report_data.get('discounts', []))} | Scope: {branch_label}"
    ws5["A2"].font = meta_font
    ws5["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws5.row_dimensions[2].height = 18

    ws5.row_dimensions[4].height = 24
    for idx, h in enumerate(dsc_headers, start=1):
        c = ws5.cell(row=4, column=idx, value=h)
        c.font = header_font
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal="center" if idx in [1, 3, 8] else ("right" if idx == 7 else "left"), vertical="center")
        c.border = THIN_BORDER

    cur_row = 5
    tot_dsc_sheet = 0.0
    for idx, d in enumerate(report_data.get('discounts', []), start=1):
        amt = float(d.get('amount', 0.0))
        tot_dsc_sheet += amt
        ws5.row_dimensions[cur_row].height = 20

        row_vals = [
            idx,
            d.get('student_name', '—'),
            d.get('roll_number', '—'),
            d.get('standard', '—'),
            d.get('branch_name', '—'),
            d.get('reason', 'Concession'),
            amt,
            d.get('granted_date', '—'),
            d.get('authorized_by', 'Staff'),
        ]

        for c_idx, val in enumerate(row_vals, start=1):
            cell = ws5.cell(row=cur_row, column=c_idx, value=val)
            cell.font = normal_font
            cell.border = THIN_BORDER
            if c_idx == 7:
                cell.number_format = '₹#,##0.00'
                cell.font = bold_font
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif c_idx in [1, 3, 8]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
                if c_idx == 2:
                    cell.font = bold_font
        cur_row += 1

    # Totals Row
    ws5.row_dimensions[cur_row].height = 24
    ws5.merge_cells(start_row=cur_row, start_column=1, end_row=cur_row, end_column=6)
    lbl = ws5.cell(row=cur_row, column=1, value=f"TOTAL DISCOUNTS AUTHORIZED ({len(report_data.get('discounts', []))} Concessions)")
    lbl.font = totals_bold
    lbl.alignment = Alignment(horizontal="right", vertical="center")

    tot_cell = ws5.cell(row=cur_row, column=7, value=tot_dsc_sheet)
    tot_cell.font = totals_bold
    tot_cell.number_format = '₹#,##0.00'
    tot_cell.alignment = Alignment(horizontal="right", vertical="center")

    for col in range(1, len(dsc_headers) + 1):
        c = ws5.cell(row=cur_row, column=col)
        c.border = TOTALS_BORDER
        c.fill = TOTALS_FILL

    for col in ws5.columns:
        max_len = max(len(str(c.value or '')) for c in col)
        col_letter = get_column_letter(col[0].column)
        ws5.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


# =============================================================================
# PDF EXPORT SERVICE
# =============================================================================

def export_user_report_pdf(report_data, meta=None):
    """
    Generates a publication-grade Landscape A4 PDF for User Activity Audit:
    - Official Executive Header with User Badge & Period metadata
    - 4-Card Executive KPI Summary Metric Ribbon
    - Part A: Fee Collections Inflow Register
    - Part B: Center Disbursements & Expenditures Register
    - Part C: Student Admissions Registered
    - Digital Audit Trail & Running Page Numbers
    """
    meta = meta or {}
    user_info = report_data.get('user') or {}
    user_name = user_info.get('name') or 'All Staff Members Combined'
    user_role = user_info.get('role_display') or 'All Roles'
    period_label = report_data.get('period', {}).get('label') or 'All Time'
    branch_label = meta.get('branch_label') or report_data.get('branch_label') or 'All Branches'
    session_label = meta.get('session_label') or report_data.get('session_label') or 'All Sessions'
    summary = report_data.get('summary', {})

    collections = report_data.get('collections', [])
    expenses = report_data.get('expenses', [])
    admissions = report_data.get('admissions', [])

    tot_col = float(summary.get('total_collected', 0.0))
    cnt_col = int(summary.get('payments_count', 0))
    tot_exp = float(summary.get('total_expenses', 0.0))
    cnt_exp = int(summary.get('expenses_count', 0))
    net_val = tot_col - tot_exp
    cnt_adm = int(summary.get('admissions_count', 0))

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=20,
        rightMargin=20,
        topMargin=20,
        bottomMargin=30,
    )

    styles = getSampleStyleSheet()

    section_title = ParagraphStyle(
        'UserRptSecTitle',
        parent=styles['Normal'],
        fontName=FONT_FAMILY_BOLD,
        fontSize=9.5,
        textColor=colors.HexColor(COLOR_NAVY_PRIMARY),
        leading=13,
    )
    cell_normal = ParagraphStyle(
        'UserRptNorm',
        parent=styles['Normal'],
        fontName=FONT_FAMILY,
        fontSize=7,
        leading=9,
        textColor=colors.HexColor(COLOR_SLATE_TEXT),
    )
    cell_bold = ParagraphStyle(
        'UserRptBld',
        parent=styles['Normal'],
        fontName=FONT_FAMILY_BOLD,
        fontSize=7,
        leading=9,
        textColor=colors.HexColor(COLOR_SLATE_TEXT),
    )
    cell_center = ParagraphStyle(
        'UserRptCenter',
        parent=styles['Normal'],
        fontName=FONT_FAMILY,
        fontSize=7,
        leading=9,
        textColor=colors.HexColor(COLOR_SLATE_TEXT),
        alignment=1,
    )
    cell_green = ParagraphStyle(
        'UserRptGreen',
        parent=styles['Normal'],
        fontName=FONT_FAMILY_BOLD,
        fontSize=7,
        leading=9,
        textColor=colors.HexColor(COLOR_GREEN_PAID),
        alignment=2,
    )
    cell_red = ParagraphStyle(
        'UserRptRed',
        parent=styles['Normal'],
        fontName=FONT_FAMILY_BOLD,
        fontSize=7,
        leading=9,
        textColor=colors.HexColor(COLOR_RED_PENDING),
        alignment=2,
    )
    hdr_style = ParagraphStyle(
        'UserRptHdr',
        parent=styles['Normal'],
        fontName=FONT_FAMILY_BOLD,
        fontSize=7.5,
        leading=9.5,
        textColor=colors.white,
        alignment=1,
    )
    hdr_right_style = ParagraphStyle(
        'UserRptHdrR',
        parent=styles['Normal'],
        fontName=FONT_FAMILY_BOLD,
        fontSize=7.5,
        leading=9.5,
        textColor=colors.white,
        alignment=2,
    )

    story = []

    # 1. Executive Header
    header_meta = {
        'branch_name': branch_label,
        'session_name': session_label,
        'date_str': f"Staff: {user_name} ({user_role}) • Period: {period_label}",
        'records_count': len(collections) + len(expenses) + len(admissions),
    }
    story.extend(build_executive_pdf_header(
        doc_title=f"STAFF OPERATIONS AUDIT: {user_name.upper()}",
        meta=header_meta,
        report_badge_label="NET CASH:",
        report_badge_value=format_currency_inr(net_val),
    ))

    # 2. Executive 4-Metric Bar (Total 802pt)
    summary_bar = [
        [
            Paragraph(f"<font size=7 color='#047857'><b>TOTAL COLLECTIONS</b></font><br/><font size=10 color='#065F46'><b>{format_currency_inr(tot_col)}</b></font><br/><font size=6.5 color='#047857'>{cnt_col} Receipts Realized</font>", styles['Normal']),
            Paragraph(f"<font size=7 color='#BE123C'><b>CENTER EXPENSES</b></font><br/><font size=10 color='#9F1239'><b>{format_currency_inr(tot_exp)}</b></font><br/><font size=6.5 color='#BE123C'>{cnt_exp} Vouchers Disbursed</font>", styles['Normal']),
            Paragraph(f"<font size=7 color='#{'047857' if net_val >= 0 else 'BE123C'}'><b>NET CASH MOVEMENT</b></font><br/><font size=10 color='#{'065F46' if net_val >= 0 else '9F1239'}'><b>{format_currency_inr(net_val)}</b></font><br/><font size=6.5 color='#64748B'>{'Surplus Contribution' if net_val >= 0 else 'Operational Deficit'}</font>", styles['Normal']),
            Paragraph(f"<font size=7 color='#1E40AF'><b>NEW ENROLLMENTS</b></font><br/><font size=10 color='#1E3A8A'><b>{cnt_adm} ADMISSIONS</b></font><br/><font size=6.5 color='#1E40AF'>Fee Vol: {format_currency_inr(summary.get('total_agreed_fees', 0.0))}</font>", styles['Normal']),
        ]
    ]
    summary_tbl = Table(summary_bar, colWidths=[200, 200, 200, 202])
    summary_tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor(COLOR_BG_CARD)),
        ('BOX', (0, 0), (-1, -1), 0.75, colors.HexColor(COLOR_BORDER)),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(summary_tbl)
    story.append(Spacer(1, 10))

    # 3. Part A: Fee Collections Inflows Table (Total width 802pt)
    # [22, 115, 60, 135, 55, 80, 80, 50, 105, 100] = 802
    story.append(Paragraph(f"<b>Part A: Fee Payments Collected by {user_name} ({len(collections)} Receipts Realized)</b>", section_title))
    story.append(Spacer(1, 4))

    col_headers_pdf = [
        Paragraph("#", hdr_style),
        Paragraph("Receipt No", hdr_style),
        Paragraph("Date", hdr_style),
        Paragraph("Student Name", hdr_style),
        Paragraph("Roll No", hdr_style),
        Paragraph("Class / Course", hdr_style),
        Paragraph("Branch", hdr_style),
        Paragraph("Mode", hdr_style),
        Paragraph("Received By", hdr_style),
        Paragraph("Amount (Rs.)", hdr_right_style),
    ]
    col_rows_pdf = [col_headers_pdf]
    for idx, p in enumerate(collections, start=1):
        amt = float(p.get('amount', 0.0))
        col_rows_pdf.append([
            Paragraph(str(idx), cell_center),
            Paragraph(str(p.get('receipt_number', '—')), cell_bold),
            Paragraph(str(p.get('payment_date', '—')), cell_center),
            Paragraph(str(p.get('student_name', '—')), cell_bold),
            Paragraph(str(p.get('roll_number', '—')), cell_center),
            Paragraph(str(p.get('standard', '—')), cell_normal),
            Paragraph(str(p.get('branch_name', '—')), cell_normal),
            Paragraph(str(p.get('payment_mode', '—')), cell_center),
            Paragraph(str(p.get('received_by', 'Staff')), cell_normal),
            Paragraph(format_currency_inr(amt), cell_green),
        ])

    col_rows_pdf.append([
        Paragraph(f"<b>TOTAL INFLOWS ({len(collections)} Receipts Realized)</b>", cell_bold),
        "", "", "", "", "", "", "", "",
        Paragraph(format_currency_inr(tot_col), cell_green),
    ])

    col_tbl = Table(col_rows_pdf, colWidths=[22, 115, 60, 135, 55, 80, 80, 50, 105, 100])
    col_tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(COLOR_NAVY_SECONDARY)),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 2.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
        ('GRID', (0, 0), (-1, -2), 0.5, colors.HexColor(COLOR_BORDER)),
        ('SPAN', (0, -1), (8, -1)),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor(COLOR_BG_TOTALS)),
        ('LINEABOVE', (0, -1), (-1, -1), 1.2, colors.HexColor(COLOR_NAVY_PRIMARY)),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#F8FAFC')]),
    ]))
    story.append(col_tbl)
    story.append(Spacer(1, 12))

    # 4. Part B: Center Expenditures Table (Total width 802pt)
    # [22, 100, 60, 110, 175, 110, 75, 50, 100] = 802
    story.append(Paragraph(f"<b>Part B: Center Expenses Recorded by {user_name} ({len(expenses)} Vouchers Incurred)</b>", section_title))
    story.append(Spacer(1, 4))

    exp_headers_pdf = [
        Paragraph("#", hdr_style),
        Paragraph("Voucher / Ref", hdr_style),
        Paragraph("Date", hdr_style),
        Paragraph("Category", hdr_style),
        Paragraph("Description / Purpose", hdr_style),
        Paragraph("Paid To / Vendor", hdr_style),
        Paragraph("Branch", hdr_style),
        Paragraph("Mode", hdr_style),
        Paragraph("Amount (Rs.)", hdr_right_style),
    ]
    exp_rows_pdf = [exp_headers_pdf]
    for idx, e in enumerate(expenses, start=1):
        amt = float(e.get('amount', 0.0))
        exp_rows_pdf.append([
            Paragraph(str(idx), cell_center),
            Paragraph(str(e.get('reference_number', '—')), cell_bold),
            Paragraph(str(e.get('expense_date', '—')), cell_center),
            Paragraph(str(e.get('category_name', 'General')), cell_bold),
            Paragraph(str(e.get('description', '—')), cell_normal),
            Paragraph(str(e.get('vendor', '—')), cell_normal),
            Paragraph(str(e.get('branch_name', '—')), cell_normal),
            Paragraph(str(e.get('payment_mode', '—')), cell_center),
            Paragraph(format_currency_inr(amt), cell_red),
        ])

    exp_rows_pdf.append([
        Paragraph(f"<b>TOTAL DISBURSEMENTS ({len(expenses)} Vouchers Incurred)</b>", cell_bold),
        "", "", "", "", "", "", "",
        Paragraph(format_currency_inr(tot_exp), cell_red),
    ])

    exp_tbl = Table(exp_rows_pdf, colWidths=[22, 100, 60, 110, 175, 110, 75, 50, 100])
    exp_tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(COLOR_NAVY_SECONDARY)),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 2.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
        ('GRID', (0, 0), (-1, -2), 0.5, colors.HexColor(COLOR_BORDER)),
        ('SPAN', (0, -1), (7, -1)),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor(COLOR_BG_TOTALS)),
        ('LINEABOVE', (0, -1), (-1, -1), 1.2, colors.HexColor(COLOR_NAVY_PRIMARY)),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#F8FAFC')]),
    ]))
    story.append(exp_tbl)
    story.append(Spacer(1, 12))

    # 5. Part C: Student Admissions Registered Table (Total width 802pt)
    # [22, 90, 50, 140, 100, 90, 60, 60, 90, 100] = 802
    story.append(Paragraph(f"<b>Part C: Student Admissions Configured by {user_name} ({len(admissions)} Enrollments)</b>", section_title))
    story.append(Spacer(1, 4))

    adm_headers_pdf = [
        Paragraph("#", hdr_style),
        Paragraph("Student Code", hdr_style),
        Paragraph("Roll No", hdr_style),
        Paragraph("Student Name", hdr_style),
        Paragraph("Class / Course", hdr_style),
        Paragraph("Branch", hdr_style),
        Paragraph("Session", hdr_style),
        Paragraph("Date", hdr_style),
        Paragraph("Agreed Fee", hdr_right_style),
        Paragraph("Net Payable (Rs.)", hdr_right_style),
    ]
    adm_rows_pdf = [adm_headers_pdf]
    tot_adm_agr = 0.0
    tot_adm_net = 0.0
    for idx, a in enumerate(admissions, start=1):
        a_fee = float(a.get('total_fees', 0.0))
        n_fee = float(a.get('net_payable', 0.0))
        tot_adm_agr += a_fee
        tot_adm_net += n_fee
        adm_rows_pdf.append([
            Paragraph(str(idx), cell_center),
            Paragraph(str(a.get('student_code', '—')), cell_bold),
            Paragraph(str(a.get('roll_number', '—')), cell_center),
            Paragraph(str(a.get('student_name', '—')), cell_bold),
            Paragraph(str(a.get('standard', '—')), cell_normal),
            Paragraph(str(a.get('branch_name', '—')), cell_normal),
            Paragraph(str(a.get('session_name', '—')), cell_center),
            Paragraph(str(a.get('admission_date', '—')), cell_center),
            Paragraph(format_currency_inr(a_fee), cell_normal),
            Paragraph(format_currency_inr(n_fee), cell_green),
        ])

    adm_rows_pdf.append([
        Paragraph(f"<b>TOTAL ENROLLMENTS ({len(admissions)} Students Enrolled)</b>", cell_bold),
        "", "", "", "", "", "", "",
        Paragraph(format_currency_inr(tot_adm_agr), cell_bold),
        Paragraph(format_currency_inr(tot_adm_net), cell_green),
    ])

    adm_tbl = Table(adm_rows_pdf, colWidths=[22, 90, 50, 140, 100, 90, 60, 60, 90, 100])
    adm_tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(COLOR_NAVY_SECONDARY)),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 2.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5),
        ('GRID', (0, 0), (-1, -2), 0.5, colors.HexColor(COLOR_BORDER)),
        ('SPAN', (0, -1), (7, -1)),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor(COLOR_BG_TOTALS)),
        ('LINEABOVE', (0, -1), (-1, -1), 1.2, colors.HexColor(COLOR_NAVY_PRIMARY)),
        ('ROWBACKGROUNDS', (0, 1), (-1, -2), [colors.white, colors.HexColor('#F8FAFC')]),
    ]))
    story.append(adm_tbl)

    doc.build(story, canvasmaker=ExecutiveLandscapeCanvas)
    buffer.seek(0)
    return buffer.getvalue()
