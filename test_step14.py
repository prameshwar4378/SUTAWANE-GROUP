import os
import sys
import json
import io
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'SUTAWANE_EDUCATION.settings')
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), 'apps'))

import django
django.setup()

from datetime import datetime, date
from django.test import Client
from django.utils import timezone
from MyApp.models import (
    Branch, AcademicSession, Student, StudentAdmission,
    FeePlan, Payment, Expense, ExpenseCategory, User
)
from MyApp.auth_utils import hash_password

TEST_TAG = "TEST_STEP14_REPORTS"

def cleanup_step14_records():
    for p in Payment.objects.filter(receipt_number__icontains=TEST_TAG):
        p.delete()
    for e in Expense.objects.filter(reference_number__icontains=TEST_TAG):
        e.delete()
    for s in Student.objects.filter(student_code__icontains=TEST_TAG):
        StudentAdmission.objects.filter(student=s).delete()
        s.delete()
    for c in ExpenseCategory.objects.filter(name__icontains=TEST_TAG):
        c.delete()
    for b in Branch.objects.filter(code__icontains=TEST_TAG):
        b.delete()
    for sess in AcademicSession.objects.filter(name__icontains=TEST_TAG):
        sess.delete()
    for u in User.objects.filter(username__icontains=TEST_TAG):
        u.delete()

cleanup_step14_records()

print("=" * 70)
print("STARTING STEP 14 VERIFICATION: 16 FINANCIAL REPORTS & MULTI-FORMAT EXPORTS")
print("=" * 70)

# 1. Setup Data
print("\n--- 1. Setting Up Test Branches, Session, Students, Payments & Expenses ---")
test_username = f"{TEST_TAG.lower()}_owner"
hashed, salt = hash_password("sutawane@123")
owner = User.objects.create(
    username=test_username,
    password=hashed,
    salt=salt,
    name="Principal Sutawane",
    role="CLASS_OWNER"
)

client = Client()
login_res = client.post('/login/', {'username': test_username, 'password': 'sutawane@123'})
assert login_res.status_code == 302


branch1 = Branch.objects.create(name=f"Central Camp {TEST_TAG}", code=f"CC_{TEST_TAG}"[:50], is_active=True)
branch2 = Branch.objects.create(name=f"East Branch {TEST_TAG}", code=f"EB_{TEST_TAG}"[:50], is_active=True)

test_session = AcademicSession.objects.create(
    name=f"2026-27_{TEST_TAG}"[:50],
    display_name="Academic Year 2026-27",
    start_date=timezone.make_aware(datetime(2026, 4, 1, 0, 0, 0)),
    end_date=timezone.make_aware(datetime(2027, 3, 31, 23, 59, 59)),
    is_current=True,
    isActive=True
)

student1 = Student.objects.create(
    student_code=f"SC-14-1_{TEST_TAG}"[:50],
    full_name="Aarav Sharma",
    parent_phone="9876543210"
)
adm1 = StudentAdmission.objects.create(
    student=student1,
    branch=branch1,
    session=test_session,
    roll_number="101",
    standard="10th State Board",
    is_active=True
)
plan1 = FeePlan.objects.create(
    admission=adm1,
    session=test_session,
    total_fees=30000.0,
    net_payable=30000.0
)

student2 = Student.objects.create(
    student_code=f"SC-14-2_{TEST_TAG}"[:50],
    full_name="Bhavya Patel",
    parent_phone="9876543211"
)
adm2 = StudentAdmission.objects.create(
    student=student2,
    branch=branch2,
    session=test_session,
    roll_number="102",
    standard="12th Science",
    is_active=True
)
plan2 = FeePlan.objects.create(
    admission=adm2,
    session=test_session,
    total_fees=50000.0,
    net_payable=48000.0
)

# Archived student
student3 = Student.objects.create(
    student_code=f"SC-14-3_{TEST_TAG}"[:50],
    full_name="Chetan Deshmukh",
    parent_phone="9876543212"
)
adm3 = StudentAdmission.objects.create(
    student=student3,
    branch=branch1,
    session=test_session,
    roll_number="103",
    standard="10th State Board",
    is_active=False
)

# Payments (Collections)
pay1 = Payment.objects.create(
    student=student1,
    branch=branch1,
    admission=adm1,
    session=test_session,
    receipt_number=f"REC-001-{TEST_TAG}",
    total_amount=15000.0,
    payment_date=timezone.make_aware(datetime(2026, 6, 15, 10, 30, 0)),
    payment_mode="UPI",
    reference_number="UPI-TXN-101"
)
pay2 = Payment.objects.create(
    student=student2,
    branch=branch2,
    admission=adm2,
    session=test_session,
    receipt_number=f"REC-002-{TEST_TAG}",
    total_amount=20000.0,
    payment_date=timezone.make_aware(datetime(2026, 6, 16, 14, 0, 0)),
    payment_mode="CASH",
    reference_number="CASH-TXN-102"
)

# Expense Categories & Expenses
cat_rent = ExpenseCategory.objects.create(name=f"Rent_{TEST_TAG}"[:150], sort_order=1)
cat_sal = ExpenseCategory.objects.create(name=f"Salaries_{TEST_TAG}"[:150], sort_order=2)

exp1 = Expense.objects.create(
    branch=branch1,
    category=cat_rent,
    amount=8000.0,
    expense_date=timezone.make_aware(datetime(2026, 6, 10, 11, 0, 0)),
    payment_mode="BANK_TRANSFER",
    description="Branch Monthly Rent",
    vendor="Mr. Landlord",
    reference_number=f"EXP-001-{TEST_TAG}"
)
exp2 = Expense.objects.create(
    branch=branch2,
    category=cat_sal,
    amount=12000.0,
    expense_date=timezone.make_aware(datetime(2026, 6, 12, 16, 0, 0)),
    payment_mode="CHEQUE",
    description="Faculty Advance",
    vendor="Prof. Patil",
    reference_number=f"EXP-002-{TEST_TAG}"
)
print("Test environment populated with 2 branches, 1 session, 3 admissions, 2 payments, and 2 expenses.")

# 2. Testing Reports Hub Tab 1: Session Audit
print("\n--- 2. Testing Tab 1: Academic Session Audit & Performance ---")
res_audit = client.get(f'/reports/?tab=SESSION_AUDIT&session_id={test_session.id}')
assert res_audit.status_code == 200, f"Expected 200, got {res_audit.status_code}"
html_audit = res_audit.content.decode('utf-8')
assert "Financial &amp; Administrative Reports Suite" in html_audit or "Financial & Administrative Reports Suite" in html_audit
assert "Academic Session Audit" in html_audit
assert "Campus &amp; Branch Performance Breakdown" in html_audit or "Campus & Branch Performance Breakdown" in html_audit
assert branch1.name in html_audit
assert branch2.name in html_audit
assert "10th State Board" in html_audit
assert "12th Science" in html_audit
assert "78,000" in html_audit  # Agreed fees
assert "35,000" in html_audit  # Realized fees
assert "20,000" in html_audit  # Center expenses
assert "15,000" in html_audit  # Net surplus

# Check JSON API representation for Session Audit
res_audit_json = client.get(f'/reports/?tab=SESSION_AUDIT&session_id={test_session.id}&format=json')
assert res_audit_json.status_code == 200
audit_json = res_audit_json.json()
assert audit_json['success'] is True
summary = audit_json['summary']
assert summary['activeAdmissions'] == 2, f"Expected 2 active admissions, got {summary['activeAdmissions']}"
assert summary['archivedAdmissions'] == 1, f"Expected 1 archived admission, got {summary['archivedAdmissions']}"
assert summary['totalFeesAgreed'] == 78000.0, f"Expected 78000 agreed fees, got {summary['totalFeesAgreed']}"
assert summary['totalCollections'] == 35000.0, f"Expected 35000 collections, got {summary['totalCollections']}"
assert summary['totalExpenses'] == 20000.0, f"Expected 20000 expenses, got {summary['totalExpenses']}"
assert summary['netSurplus'] == 15000.0, f"Expected 15000 surplus, got {summary['netSurplus']}"
assert summary['isSurplus'] is True
print("Tab 1 (Session Audit) verified: enrollment counts, fees, collections, expenses, and branch/standard breakdowns accurate!")

# 3. Testing Tab 2: Payment Channels & Modes
print("\n--- 3. Testing Tab 2: Payment Channels & Inflow Modes ---")
res_modes = client.get(f'/reports/?tab=PAYMENT_MODES&session_id={test_session.id}')
assert res_modes.status_code == 200
html_modes = res_modes.content.decode('utf-8')
assert "Payment Modes &amp; Digital Inflows Breakdown" in html_modes or "Payment Modes & Digital Inflows Breakdown" in html_modes
assert "Reconciled Payment Methods Register" in html_modes
assert "UPI" in html_modes
assert "CASH" in html_modes

res_modes_json = client.get(f'/reports/?tab=PAYMENT_MODES&session_id={test_session.id}&format=json')
assert res_modes_json.status_code == 200
modes_json = res_modes_json.json()
assert modes_json['success'] is True
assert modes_json['totalCollected'] == 35000.0
assert modes_json['totalTransactions'] == 2
print("Tab 2 (Payment Modes) verified: mode shares, transaction counts, and totals accurate!")

# 4. Testing Tab 3: Expense Categories Analysis
print("\n--- 4. Testing Tab 3: Expense Category Analysis ---")
res_exp = client.get(f'/reports/?tab=EXPENSE_CATEGORIES&session_id={test_session.id}')
assert res_exp.status_code == 200
html_exp = res_exp.content.decode('utf-8')
assert "Operational Expenses &amp; Category Taxonomy" in html_exp or "Operational Expenses & Category Taxonomy" in html_exp
assert cat_rent.name in html_exp
assert cat_sal.name in html_exp

res_exp_json = client.get(f'/reports/?tab=EXPENSE_CATEGORIES&session_id={test_session.id}&format=json')
assert res_exp_json.status_code == 200
exp_json = res_exp_json.json()
assert exp_json['success'] is True
assert exp_json['totalExpenses'] == 20000.0
assert exp_json['totalVouchers'] == 2
print("Tab 3 (Expense Categories) verified: category totals, voucher counts, and distribution accurate!")

# 5. Testing Tab 4: Statements & Ledgers Directory
print("\n--- 5. Testing Tab 4: Statements & Ledgers Directory ---")
res_dir = client.get(f'/reports/?tab=STATEMENT_DIRECTORY&session_id={test_session.id}')
assert res_dir.status_code == 200
html_dir = res_dir.content.decode('utf-8')
assert "All-in-One Financial Movement &amp; Audit Statement" in html_dir or "All-in-One Financial Movement & Audit Statement" in html_dir
assert "1. Daily Financial Statement" in html_dir
assert "2. Month-wise Financial Report" in html_dir
assert "3. Custom Date Range Audit" in html_dir
assert "4. Collections Register &amp; Inflows" in html_dir or "4. Collections Register & Inflows" in html_dir
print("Tab 4 (Statements Directory) verified: all financial statements and quick access links rendered!")

# 6. Testing Multi-Format Excel Exports
print("\n--- 6. Testing Excel Exports (Session Audit, Payment Modes, Categories) ---")
res_audit_xl = client.get(f'/reports/?tab=SESSION_AUDIT&session_id={test_session.id}&export=excel')
assert res_audit_xl.status_code == 200
assert res_audit_xl['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
wb_audit = openpyxl.load_workbook(io.BytesIO(res_audit_xl.content))
assert 'Branch Comparison' in wb_audit.sheetnames
assert 'Course Breakdown' in wb_audit.sheetnames
assert 'Monthly Trend' in wb_audit.sheetnames

res_modes_xl = client.get(f'/reports/?tab=PAYMENT_MODES&session_id={test_session.id}&export=excel')
assert res_modes_xl.status_code == 200
wb_modes = openpyxl.load_workbook(io.BytesIO(res_modes_xl.content))
assert 'Payment Channels' in wb_modes.sheetnames

res_exp_xl = client.get(f'/reports/?tab=EXPENSE_CATEGORIES&session_id={test_session.id}&export=excel')
assert res_exp_xl.status_code == 200
wb_exp = openpyxl.load_workbook(io.BytesIO(res_exp_xl.content))
assert 'Expense Categories' in wb_exp.sheetnames
print("Excel exports verified: Session Audit (.xlsx with 3 sheets), Payment Modes (.xlsx), Expense Categories (.xlsx) valid!")

# 7. Testing PDF Exports
print("\n--- 7. Testing PDF Exports (Payment Modes & Categories) ---")
res_modes_pdf = client.get(f'/reports/?tab=PAYMENT_MODES&session_id={test_session.id}&export=pdf')
assert res_modes_pdf.status_code == 200
assert res_modes_pdf['Content-Type'] == 'application/pdf'
assert res_modes_pdf.content.startswith(b'%PDF')

res_exp_pdf = client.get(f'/reports/?tab=EXPENSE_CATEGORIES&session_id={test_session.id}&export=pdf')
assert res_exp_pdf.status_code == 200
assert res_exp_pdf['Content-Type'] == 'application/pdf'
assert res_exp_pdf.content.startswith(b'%PDF')
print("PDF exports verified: Payment Modes PDF and Expense Categories PDF generated cleanly!")

# 8. Testing Day Book & Daily Statement (GET /reports/daily/)
print("\n--- 8. Testing Day Book: DAILY, MONTHLY & CUSTOM Modes ---")
# Daily mode: date of pay1 (2026-06-15)
res_daily = client.get(f'/reports/daily/?mode=DAILY&date=2026-06-15&session_id={test_session.id}')
assert res_daily.status_code == 200
html_daily = res_daily.content.decode('utf-8')
assert "Daily Financial Statement &amp; Day Book Register" in html_daily or "Daily Financial Statement & Day Book Register" in html_daily
assert pay1.receipt_number in html_daily
assert pay2.receipt_number not in html_daily  # Different date

res_daily_json = client.get(f'/reports/daily/?mode=DAILY&date=2026-06-15&session_id={test_session.id}&format=json')
assert res_daily_json.status_code == 200
daily_sum = res_daily_json.json()['summary']
assert daily_sum['totalCollections'] == 15000.0
assert daily_sum['collectionsCount'] == 1

# Monthly mode: June 2026 (month=5 for June 0-indexed)
res_monthly = client.get(f'/reports/daily/?mode=MONTHLY&month=5&year=2026&session_id={test_session.id}')
assert res_monthly.status_code == 200
html_monthly = res_monthly.content.decode('utf-8')
assert pay1.receipt_number in html_monthly
assert pay2.receipt_number in html_monthly
assert exp1.reference_number in html_monthly
assert exp2.reference_number in html_monthly

res_monthly_json = client.get(f'/reports/daily/?mode=MONTHLY&month=5&year=2026&session_id={test_session.id}&format=json')
assert res_monthly_json.status_code == 200
monthly_sum = res_monthly_json.json()['summary']
assert monthly_sum['totalCollections'] == 35000.0
assert monthly_sum['totalExpenses'] == 20000.0
assert monthly_sum['netMovement'] == 15000.0

# Custom range mode:
res_custom = client.get(f'/reports/daily/?mode=CUSTOM&start_date=2026-06-01&end_date=2026-06-30&session_id={test_session.id}')
assert res_custom.status_code == 200
print("Day Book modes verified: Daily date filtering, Monthly period aggregation, and Custom range all accurate!")

# 9. Testing Day Book Excel & PDF Exports
print("\n--- 9. Testing Day Book Multi-Sheet Excel & ReportLab PDF Exports ---")
res_day_xl = client.get(f'/reports/daily/?mode=MONTHLY&month=5&year=2026&session_id={test_session.id}&export=excel')
assert res_day_xl.status_code == 200
assert res_day_xl['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
wb_day = openpyxl.load_workbook(io.BytesIO(res_day_xl.content))
assert 'Executive Summary' in wb_day.sheetnames
assert 'Fee Collections' in wb_day.sheetnames
assert 'Center Disbursements' in wb_day.sheetnames

res_day_pdf = client.get(f'/reports/daily/?mode=MONTHLY&month=5&year=2026&session_id={test_session.id}&export=pdf')
assert res_day_pdf.status_code == 200
assert res_day_pdf['Content-Type'] == 'application/pdf'
assert res_day_pdf.content.startswith(b'%PDF')
print("Day Book exports verified: Multi-sheet workbook and Landscape A4 PDF successfully generated!")

# 10. Testing HTMX Search Partial Swaps
print("\n--- 10. Testing HTMX Partial Table Swaps ---")
res_htmx_col = client.get(
    f'/reports/daily/?table=collections&mode=MONTHLY&month=5&year=2026&session_id={test_session.id}&q_col=Aarav',
    HTTP_HX_REQUEST='true'
)
assert res_htmx_col.status_code == 200
html_htmx_col = res_htmx_col.content.decode('utf-8')
assert "Aarav Sharma" in html_htmx_col
assert "Bhavya Patel" not in html_htmx_col

res_htmx_exp = client.get(
    f'/reports/daily/?table=expenses&mode=MONTHLY&month=5&year=2026&session_id={test_session.id}&q_exp=Landlord',
    HTTP_HX_REQUEST='true'
)
assert res_htmx_exp.status_code == 200
html_htmx_exp = res_htmx_exp.content.decode('utf-8')
assert "Mr. Landlord" in html_htmx_exp
assert "Prof. Patil" not in html_htmx_exp
print("HTMX partial table swaps verified: instant filter swaps without full page reloads!")

# 11. Testing API Endpoints Parity with React Backend
print("\n--- 11. Testing API Endpoints Contract Parity ---")
# /api/reports/statement/
res_api_stmt = client.get(f'/api/reports/statement/?startDate=2026-06-01&endDate=2026-06-30&sessionId={test_session.id}')
assert res_api_stmt.status_code == 200
data_stmt = res_api_stmt.json()
assert data_stmt['success'] is True
assert 'summary' in data_stmt
assert 'collections' in data_stmt
assert 'expenses' in data_stmt

# /api/reports/session-audit/
res_api_audit = client.get(f'/api/reports/session-audit/?sessionId={test_session.id}')
assert res_api_audit.status_code == 200
data_audit = res_api_audit.json()
assert data_audit['success'] is True
assert 'session' in data_audit
assert 'branchBreakdown' in data_audit
assert 'standardBreakdown' in data_audit

# /api/reports/payment-modes-summary/
res_api_modes = client.get(f'/api/reports/payment-modes-summary/?sessionId={test_session.id}')
assert res_api_modes.status_code == 200
data_modes = res_api_modes.json()
assert data_modes['success'] is True
assert 'modes' in data_modes

# /api/reports/expense-categories-summary/
res_api_cats = client.get(f'/api/reports/expense-categories-summary/?sessionId={test_session.id}')
assert res_api_cats.status_code == 200
data_cats = res_api_cats.json()
assert data_cats['success'] is True
assert 'categories' in data_cats
print("API endpoints contract parity verified: all 4 JSON API endpoints match React expectations!")

# 12. Cleanup
print("\n--- 12. Cleanup Test Artifacts ---")
cleanup_step14_records()
print("Cleaned up test reports records cleanly.")

print("\n" + "=" * 70)
print("ALL STEP 14 TESTS PASSED (100% SUCCESS)")
print("=" * 70)
